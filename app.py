from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
from uuid import uuid4
from urllib.parse import parse_qs, quote, urlparse
import re
from html import escape
import base64
import inspect
import math
import zipfile
from zoneinfo import ZoneInfo
from functools import partial
import io
import time
from concurrent.futures import ThreadPoolExecutor, wait as wait_futures
import json
import unicodedata

import pandas as pd
import requests
import streamlit as st
import folium
from streamlit_folium import st_folium
from supabase import Client, create_client


st.set_page_config(page_title="Ofis Portföyü", page_icon="🏠", layout="wide")
PHOTO_BUCKET = "property-photos"
_TAB_PARAMS = inspect.signature(st.tabs).parameters
# Yeni Streamlit sürümlerinde sekmeler session_state ile kodla değiştirilebilir.
TABS_STATE_SUPPORTED = "key" in _TAB_PARAMS and "on_change" in _TAB_PARAMS
MAX_PHOTO_BYTES = 10 * 1024 * 1024
DOWNLOAD_CALLABLE = "callable" in (st.download_button.__doc__ or "").lower()
TZ = ZoneInfo("Europe/Istanbul")
try:
    from streamlit.delta_generator import DeltaGenerator as _DG

    if "placeholder" in inspect.signature(_DG.multiselect).parameters and not getattr(_DG.multiselect, "_tr_patched", False):
        _multiselect_original = _DG.multiselect

        def _multiselect_tr(self, *args, **kwargs):
            kwargs.setdefault("placeholder", "Seçiniz")
            return _multiselect_original(self, *args, **kwargs)

        _multiselect_tr._tr_patched = True
        _DG.multiselect = _multiselect_tr
        st.multiselect = lambda *args, **kwargs: _multiselect_tr(st._main, *args, **kwargs)
except Exception:
    pass
DEFAULT_OFFICE_NAME = "Ofis Portföyü"
DEFAULT_HOME_TITLE = "Ofis genel görünümü"
NEW_SQL_FILE = "supabase_yeni_ozellikler.sql"
OFFER_STATUSES = ["Beklemede", "Karşı teklif", "Kabul edildi", "Reddedildi", "İptal"]
DEFAULT_PROPERTY_IMAGE = Path(__file__).parent / "assets" / "no-property-photo.svg"

st.markdown("""
<style>
.block-container { padding-top: 4.2rem; padding-bottom: 5rem; max-width: 1320px; }
div[data-testid="stVerticalBlock"] { gap: .6rem; }
button[role="tab"] { padding-top: .35rem; padding-bottom: .35rem; }
@media (max-width: 760px) {
  .block-container { padding: 3.8rem .8rem 4.5rem; }
  div[data-testid="stHorizontalBlock"] { flex-direction: column !important; gap: .45rem !important; }
  div[data-testid="stHorizontalBlock"] > div { width: 100% !important; flex: 1 1 100% !important; }
  [class*="st-key-photo_grid"] div[data-testid="stHorizontalBlock"] { flex-direction: row !important; flex-wrap: wrap !important; }
  [class*="st-key-photo_grid"] div[data-testid="stHorizontalBlock"] > div { width: calc(33% - .4rem) !important; flex: 0 0 calc(33% - .4rem) !important; }
  [class*="st-key-listing_row_"] div[data-testid="stHorizontalBlock"] { flex-direction: row !important; flex-wrap: nowrap !important; gap: .4rem !important; align-items: center !important; }
  [class*="st-key-listing_row_"] div[data-testid="stHorizontalBlock"] > div:first-child { width: 64px !important; flex: 0 0 64px !important; min-width: 64px !important; }
  [class*="st-key-listing_row_"] div[data-testid="stHorizontalBlock"] > div:last-child { width: auto !important; flex: 1 1 auto !important; min-width: 0 !important; }
  h1 { font-size: 1.5rem !important; }
  h2 { font-size: 1.25rem !important; }
}
.st-key-topbar { background: #ffffff; border: 1px solid #e4e7ec; border-radius: 16px; padding: .5rem 1rem; box-shadow: 0 2px 12px rgba(16,24,40,.06); }
.st-key-topbar button { background: none !important; border: 0 !important; box-shadow: none !important; min-height: 0 !important; padding: .3rem .6rem !important; }
.st-key-topbar button, .st-key-topbar button * { color: #334155 !important; }
.st-key-topbar button p { font-weight: 600; font-size: .95rem; margin: 0; }
.st-key-topbar button:hover, .st-key-topbar button:hover * { color: #0284c7 !important; }
.st-key-brand { position: relative; }
.st-key-brand .brand { display: flex; align-items: center; gap: .55rem; padding: .15rem 0; }
.st-key-brand .brand-logo { font-size: 1.7rem; line-height: 1; }
.st-key-brand .brand-text { font-size: 1.3rem; font-weight: 800; color: #0f172a !important; line-height: 1.1; }
.st-key-brand .brand-text b { color: #0284c7 !important; }
.st-key-home_title { position: absolute; inset: 0; z-index: 5; }
.st-key-home_title button { width: 100% !important; height: 100% !important; opacity: 0; cursor: pointer; }
@media (max-width: 760px) {
  .st-key-topbar div[data-testid="stHorizontalBlock"] { flex-direction: row !important; flex-wrap: wrap !important; gap: .15rem !important; align-items: center !important; }
  .st-key-topbar div[data-testid="stHorizontalBlock"] > div { width: auto !important; flex: 0 0 auto !important; min-width: 0 !important; }
  .st-key-topbar div[data-testid="stHorizontalBlock"] > div:first-child { flex: 1 1 100% !important; }
  .st-key-brand .brand-text { font-size: 1.1rem; }
}
.st-key-agent_card { position: relative; }
.agent-card { background: #ffffff; border: 1px solid #e4e7ec; border-radius: 16px; padding: 1rem .8rem; text-align: center; box-shadow: 0 2px 10px rgba(16,24,40,.05); }
.agent-card .agent-img { width: 88px; height: 88px; border-radius: 50%; object-fit: cover; display: block; margin: 0 auto .55rem; border: 3px solid #ffffff; box-shadow: 0 2px 10px rgba(0,0,0,.2); }
.agent-card .agent-img.ph { background: #e0e7ff; display: flex; align-items: center; justify-content: center; font-size: 2.1rem; font-weight: 700; color: #1e3a8a !important; }
.agent-card .agent-name { font-size: 1.15rem; font-weight: 800; color: #0f172a !important; }
.agent-card .agent-role { font-size: .82rem; color: #64748b !important; margin-top: .1rem; }
.st-key-agent_open { position: absolute; inset: 0; z-index: 5; }
.st-key-agent_open button { width: 100% !important; height: 100% !important; opacity: 0; cursor: pointer; }
.st-key-agent_call a { background: #ff4d4f !important; border: 0 !important; border-radius: 12px !important; font-weight: 700; }
.st-key-agent_call a, .st-key-agent_call a * { color: #ffffff !important; }
.info-tiles { display: grid; grid-template-columns: repeat(auto-fit, minmax(118px, 1fr)); gap: .6rem; margin-top: .7rem; }
.info-tile { background: #f9fafb; border: 1px solid #e4e7ec; border-radius: 12px; padding: .7rem .4rem; text-align: center; }
.info-tile .it-icon { font-size: 1.4rem; line-height: 1.2; }
.info-tile .it-label { font-weight: 700; font-size: .84rem; color: #0f172a !important; margin-top: .2rem; }
.info-tile .it-value { font-size: .82rem; color: #64748b !important; margin-top: .1rem; }
[class*="st-key-listing_row_"] button { text-align: left; min-height: 46px; padding: .35rem .8rem; white-space: normal; font-size: .9rem; }
.summary-strip { display: flex; gap: .4rem; flex-wrap: wrap; justify-content: flex-start; margin: .5rem 0 .2rem; }
.summary-strip .chip { background: #fff; border: 1px solid #e4e7ec; border-radius: 999px; padding: .2rem .75rem; font-size: .82rem; color: #475467 !important; }
.summary-strip .chip b { color: #101828 !important; font-size: 1rem; margin-left: .35rem; }
@media (max-width: 760px) { .summary-strip { justify-content: flex-start; } }
.price-card { background: #fff; border: 1px solid #e4e7ec; border-radius: 12px; padding: .6rem .9rem; line-height: 1.45; font-size: .9rem; color: #101828 !important; }
.price-card .price { font-size: 1.7rem; font-weight: 800; color: #1d4ed8; line-height: 1.2; }
.deal-card { text-align: center; background: linear-gradient(135deg, #1e3a8a, #2563eb); border-radius: 14px; padding: .85rem 1rem; box-shadow: 0 4px 14px rgba(30,58,138,.25); }
.deal-card, .deal-card * { color: #ffffff !important; }
.deal-card .dc-top { display: flex; flex-direction: column; align-items: center; gap: .4rem; }
.deal-card .dc-price { font-size: 1.65rem; font-weight: 800; line-height: 1.15; }
.deal-card .dc-price.muted { font-size: 1.05rem; font-weight: 600; opacity: .85; }
.deal-card .dc-status { font-size: .74rem; font-weight: 700; padding: .2rem .7rem; border-radius: 999px; white-space: nowrap; }
.deal-card .dc-chips { display: flex; flex-wrap: wrap; justify-content: center; gap: .35rem; margin: .6rem 0 .5rem; }
.deal-card .dc-chip { background: rgba(255,255,255,.18); border-radius: 999px; padding: .15rem .7rem; font-size: .8rem; }
.deal-card .dc-row { font-size: .88rem; margin-top: .22rem; }
.hero { position: relative; height: 230px; border-radius: 16px; overflow: hidden; background-size: cover; background-position: center; background-color: #1e293b; }
.hero .hero-shade { position: absolute; inset: 0; background: linear-gradient(to top, rgba(8,16,36,.92) 0%, rgba(8,16,36,.45) 42%, rgba(0,0,0,0) 72%); }
.hero .hero-body { position: absolute; left: 0; right: 0; bottom: 0; padding: 1rem 1.2rem; }
.hero, .hero * { color: #ffffff !important; }
.hero .hero-badge { display: inline-block; background: #0284c7; font-size: .72rem; font-weight: 700; letter-spacing: .03em; padding: .2rem .7rem; border-radius: 999px; }
.hero .hero-title { font-size: 1.35rem; font-weight: 700; line-height: 1.25; margin: .4rem 0 .35rem; text-shadow: 0 1px 4px rgba(0,0,0,.5); }
.hero .hero-meta { font-size: .9rem; opacity: .95; }
.hero .hero-price { color: #38bdf8 !important; font-weight: 800; font-size: 1.1rem; margin-right: .6rem; }
.hero-dots { text-align: center; letter-spacing: .25rem; font-size: .8rem; color: #667085; }
.avatar-ph { width: 96px; height: 96px; border-radius: 50%; background: #e0e7ff; color: #1e3a8a !important; display: flex; align-items: center; justify-content: center; font-size: 2.2rem; font-weight: 700; }
@media (max-width: 760px) {
  .hero { height: 190px; }
  .hero .hero-title { font-size: 1.1rem; }
  [class*="st-key-hero_controls"] div[data-testid="stHorizontalBlock"] { flex-direction: row !important; flex-wrap: nowrap !important; align-items: center !important; }
  [class*="st-key-hero_controls"] div[data-testid="stHorizontalBlock"] > div { width: auto !important; flex: 1 1 auto !important; min-width: 0 !important; }
  [class*="st-key-hero_controls"] div[data-testid="stHorizontalBlock"] > div:first-child, [class*="st-key-hero_controls"] div[data-testid="stHorizontalBlock"] > div:last-child { flex: 0 0 56px !important; }
}
.property-feature-table { width: 100%; border-collapse: collapse; font-size: .84rem; }
.property-feature-table th, .property-feature-table td { padding: .28rem .55rem; border-bottom: 1px solid #eaecf0; text-align: left; background: #fff; color: #101828 !important; }
.property-feature-table th { width: 45%; color: #475467 !important; font-weight: 600; background: #f9fafb; }
</style>
""", unsafe_allow_html=True)


def get_client() -> Client:
    return create_client(st.secrets["supabase"]["url"], st.secrets["supabase"]["anon_key"])


def get_admin_client() -> Client:
    key = st.secrets["supabase"].get("service_role_key") or st.secrets["supabase"].get("secret_key")
    if not key:
        raise RuntimeError("Yönetici işlemleri için Supabase service_role/secret key ayarlanmamış.")
    return create_client(st.secrets["supabase"]["url"], key)


def now_tr() -> datetime:
    return datetime.now(TZ)


def feature_setup_message(feature: str, exc: Exception | None = None) -> None:
    st.warning(f"{feature} için Supabase'de ek tablolar gerekli. `{NEW_SQL_FILE}` dosyasını Supabase SQL Editor'da bir kez çalıştırın.")
    path = Path(__file__).parent / NEW_SQL_FILE
    if path.exists():
        with st.expander("Çalıştırılacak SQL"):
            st.code(path.read_text(encoding="utf-8"), language="sql")
    if exc is not None:
        st.caption(str(exc))


@st.cache_data(ttl=60, show_spinner=False)
def get_settings(_client: Client) -> dict:
    try:
        data = _client.table("app_settings").select("*").execute().data or []
    except Exception:
        return {}
    return {item["key"]: item.get("value") for item in data}


def load_announcements(client: Client, only_active: bool = True) -> list[dict]:
    try:
        query = client.table("announcements").select("*").order("created_at", desc=True)
        if only_active:
            query = query.eq("active", True)
        return query.execute().data or []
    except Exception:
        return []


def show_announcements(client: Client) -> None:
    items = load_announcements(client)
    for item in items[:3]:
        body = f"\n\n{item['body']}" if item.get("body") else ""
        st.info(f"**{item['title']}**{body}", icon="📢")
    if len(items) > 3:
        with st.expander(f"Diğer {len(items) - 3} duyuru"):
            for item in items[3:]:
                st.markdown(f"**{item['title']}**  \n{item.get('body') or ''}")


def settings_admin_tab(client: Client, admin_client: Client) -> None:
    settings = get_settings(client)
    st.caption("Burada yazdıklarınız tüm kullanıcılarda görünür. Ofis adı sol menüde ve PDF/Excel çıktılarının başlığında kullanılır.")
    office = st.text_input("Ofis / uygulama adı", value=settings.get("office_name") or DEFAULT_OFFICE_NAME, key="admin_office_name")
    home = st.text_input("Ana sayfa başlığı (üstteki buton)", value=settings.get("home_title") or DEFAULT_HOME_TITLE, key="admin_home_title")
    if st.button("Başlıkları kaydet", type="primary", key="admin_save_titles"):
        try:
            admin_client.table("app_settings").upsert([
                {"key": "office_name", "value": office.strip() or DEFAULT_OFFICE_NAME},
                {"key": "home_title", "value": home.strip() or DEFAULT_HOME_TITLE},
            ]).execute()
            get_settings.clear()
            st.success("Başlıklar kaydedildi.")
            st.rerun()
        except Exception as exc:
            feature_setup_message("Başlık ayarları", exc)


def announcements_admin_tab(admin_client: Client) -> None:
    try:
        items = admin_client.table("announcements").select("*").order("created_at", desc=True).execute().data or []
    except Exception as exc:
        feature_setup_message("Duyurular", exc)
        return
    generation = st.session_state.get("news_generation", 0)

    def bump() -> None:
        st.session_state["news_generation"] = generation + 1

    with st.expander("➕ Yeni duyuru ekle", expanded=not items):
        title = st.text_input("Duyuru başlığı *", key=f"news_title_{generation}")
        body = st.text_area("Duyuru metni", height=90, key=f"news_body_{generation}")
        if st.button("Duyuruyu yayınla", type="primary", key=f"news_add_{generation}"):
            if not title.strip():
                st.error("Duyuru başlığını yazın.")
            else:
                admin_client.table("announcements").insert({"title": title.strip(), "body": body.strip(), "active": True}).execute()
                bump()
                st.rerun()
    for item in items:
        item_id = str(item["id"])
        created = pd.to_datetime(item.get("created_at"), errors="coerce")
        stamp = created.strftime("%d.%m.%Y") if pd.notna(created) else ""
        with st.expander(f"{'🟢' if item.get('active') else '⚪'} {item['title']} · {stamp}"):
            new_title = st.text_input("Başlık", value=item["title"], key=f"news_edit_title_{item_id}_{generation}")
            new_body = st.text_area("Metin", value=item.get("body") or "", height=90, key=f"news_edit_body_{item_id}_{generation}")
            active = st.checkbox("Yayında (kullanıcılara gösterilsin)", value=bool(item.get("active")), key=f"news_edit_active_{item_id}_{generation}")
            b1, b2, b3 = st.columns([1, 1, 1], vertical_alignment="center")
            if b1.button("Güncelle", type="primary", key=f"news_update_{item_id}_{generation}", use_container_width=True):
                admin_client.table("announcements").update({"title": new_title.strip(), "body": new_body.strip(), "active": active, "updated_at": now_tr().isoformat()}).eq("id", item_id).execute()
                bump()
                st.rerun()
            confirm = b2.checkbox("Silmeyi onaylıyorum", key=f"news_confirm_{item_id}_{generation}")
            if b3.button("Duyuruyu sil", disabled=not confirm, key=f"news_delete_{item_id}_{generation}", use_container_width=True):
                admin_client.table("announcements").delete().eq("id", item_id).execute()
                bump()
                st.rerun()


def service_client(fallback: Client) -> Client:
    try:
        return get_admin_client()
    except Exception:
        return fallback


def login(client: Client) -> None:
    st.title("🏠 Ofis Portföyü")
    st.caption("Daireleri ve müşteri taleplerini tek ekranda yönetin.")
    with st.form("login"):
        email = st.text_input("E-posta")
        password = st.text_input("Şifre", type="password")
        submitted = st.form_submit_button("Giriş yap", type="primary")
    if submitted:
        try:
            result = client.auth.sign_in_with_password({"email": email, "password": password})
            st.session_state.user = result.user
            st.session_state.supabase_access_token = result.session.access_token
            st.session_state.supabase_refresh_token = result.session.refresh_token
            st.rerun()
        except Exception:
            st.error("Giriş yapılamadı. E-posta ve şifrenizi kontrol edin.")


def rows(client: Client, table: str) -> list[dict]:
    return client.table(table).select("*").order("created_at", desc=True).execute().data


@st.cache_data(ttl=86400, show_spinner=False)
def get_location_data() -> tuple[list[str], dict[str, list[str]]]:
    """Load the 81 provinces and their districts, cache for one day."""
    province_response = requests.get("https://api.turkiyeapi.dev/v2/datasets/2025/provinces.json", timeout=12)
    province_response.raise_for_status()
    district_response = requests.get("https://api.turkiyeapi.dev/v2/datasets/2025/districts.json", timeout=12)
    district_response.raise_for_status()
    provinces_raw = province_response.json()
    districts_raw = district_response.json()
    provinces_raw = provinces_raw.get("data", provinces_raw) if isinstance(provinces_raw, dict) else provinces_raw
    districts_raw = districts_raw.get("data", districts_raw) if isinstance(districts_raw, dict) else districts_raw
    provinces = {}
    for item in provinces_raw:
        name = item.get("name")
        if isinstance(name, dict):
            name = name.get("local") or name.get("tr")
        if name:
            provinces[str(item.get("id"))] = str(name)
    districts: dict[str, list[str]] = {name: [] for name in provinces.values()}
    for item in districts_raw:
        province_id = item.get("provinceId", item.get("province_id"))
        if isinstance(item.get("province"), dict):
            province_id = item["province"].get("id", province_id)
        province_name = provinces.get(str(province_id))
        name = item.get("name")
        if isinstance(name, dict):
            name = name.get("local") or name.get("tr")
        if province_name and name:
            districts[province_name].append(str(name))
    if len(provinces) != 81:
        raise ValueError("İl listesi beklenen 81 il sayısında değil.")
    return sorted(provinces.values(), key=str.casefold), {k: sorted(v, key=str.casefold) for k, v in districts.items()}


def parse_google_maps_location(url: str) -> tuple[str | None, str | None]:
    patterns = [r"@(-?\d+(?:\.\d+)?),(-?\d+(?:\.\d+)?)", r"!3d(-?\d+(?:\.\d+)?)!4d(-?\d+(?:\.\d+)?)"]
    for pattern in patterns:
        match = re.search(pattern, url)
        if match:
            return match.group(1), match.group(2)
    query = parse_qs(urlparse(url).query)
    for key in ("q", "query", "ll", "center"):
        value = query.get(key, [""])[0]
        match = re.search(r"(-?\d+(?:\.\d+)?),\s*(-?\d+(?:\.\d+)?)", value)
        if match:
            return match.group(1), match.group(2)
    return None, None


def upload_property_photos(client: Client, property_id: str, files: list) -> tuple[int, list[str]]:
    uploaded = 0
    errors = []
    bucket = client.storage.from_(PHOTO_BUCKET)
    for image in files:
        content = image.getvalue()
        if image.type not in {"image/jpeg", "image/png", "image/webp"}:
            errors.append(f"{image.name}: JPG, PNG veya WebP yükleyin.")
            continue
        if len(content) > MAX_PHOTO_BYTES:
            errors.append(f"{image.name}: dosya 10 MB sınırını aşıyor.")
            continue
        suffix = Path(image.name).suffix.lower()
        object_path = f"{property_id}/{uuid4().hex}{suffix}"
        try:
            bucket.upload(object_path, content, file_options={"content-type": image.type, "upsert": "false"})
            uploaded += 1
        except Exception as exc:
            errors.append(f"{image.name}: {exc}")
    return uploaded, errors


def show_property_gallery(client: Client, property_id: str) -> None:
    try:
        bucket = client.storage.from_(PHOTO_BUCKET)
        objects = bucket.list(property_id)
        if not objects:
            st.info("Bu ilana henüz fotoğraf eklenmemiş.")
            return
        columns = st.columns(3)
        for i, obj in enumerate(objects):
            path = f"{property_id}/{obj['name']}"
            response = bucket.create_signed_url(path, 3600)
            url = response.get("signedURL") or response.get("signedUrl")
            if url:
                columns[i % 3].image(url, caption=obj["name"], width="stretch")
    except Exception as exc:
        st.error("Fotoğraflar alınamadı. Supabase Storage bucket ve erişim kurulumunu kontrol edin.")
        st.caption(str(exc))


def get_property_photos(client: Client, property_id: str) -> list[dict]:
    bucket = client.storage.from_(PHOTO_BUCKET)
    objects = bucket.list(property_id, {"limit": 100, "sortBy": {"column": "name", "order": "asc"}})
    photos = []
    for obj in objects or []:
        if not obj.get("id"):
            continue
        path = f"{property_id}/{obj['name']}"
        response = bucket.create_signed_url(path, 3600)
        url = response.get("signedURL") or response.get("signedUrl")
        if url:
            photos.append({"name": obj["name"], "url": url})
    return photos


def embed_html(html: str, height: int) -> None:
    if hasattr(st, "iframe"):
        st.iframe(html, height=height)
    else:
        st.components.v1.html(html, height=height)


def embed_page(url: str, height: int) -> None:
    if hasattr(st, "iframe"):
        st.iframe(url, height=height)
    else:
        st.components.v1.iframe(url, height=height, scrolling=False)


GALLERY_HTML = """<!doctype html>
<html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<style>
*{box-sizing:border-box;-webkit-tap-highlight-color:transparent}
body{margin:0;font-family:system-ui,-apple-system,sans-serif;background:transparent}
.main{position:relative;height:290px;background:#101828;border-radius:12px;overflow:hidden;display:flex;align-items:center;justify-content:center}
.main img{max-width:100%;max-height:100%;object-fit:contain;cursor:zoom-in;display:block}
.nav{position:absolute;top:50%;margin-top:-22px;width:44px;height:44px;border:0;border-radius:50%;background:rgba(255,255,255,.88);font-size:28px;line-height:1;color:#101828;cursor:pointer}
.prev{left:8px}.next{right:8px}
.count{position:absolute;bottom:8px;right:10px;background:rgba(0,0,0,.6);color:#fff;font-size:12px;padding:3px 9px;border-radius:12px}
.thumbs{display:flex;gap:6px;margin-top:8px;overflow-x:auto;padding-bottom:4px}
.thumbs img{height:54px;width:74px;object-fit:cover;border-radius:6px;border:2px solid transparent;opacity:.65;flex:none;cursor:pointer}
.thumbs img.active{border-color:#2563eb;opacity:1}
.hint{font-size:12px;color:#667085;margin-top:2px}
</style></head><body>
<div class="main" id="main"><img id="mainImg" alt="İlan fotoğrafı"><button type="button" class="nav prev" id="prev">&#8249;</button><button type="button" class="nav next" id="next">&#8250;</button><div class="count" id="count"></div></div>
<div class="thumbs" id="thumbs"></div>
<div class="hint">Büyütmek için fotoğrafa dokunun · kaydırarak ileri geri gidebilirsiniz</div>
<script>
const photos = __PHOTOS__;
let idx = 0;
const $ = id => document.getElementById(id);
const mainImg = $("mainImg"), count = $("count"), thumbs = $("thumbs");
if (photos.length < 2) { $("prev").style.display = "none"; $("next").style.display = "none"; thumbs.style.display = "none"; }
photos.forEach((url, i) => { const t = document.createElement("img"); t.src = url; t.loading = "lazy"; t.onclick = () => show(i); thumbs.appendChild(t); });
function show(i) {
  idx = (i + photos.length) % photos.length;
  mainImg.src = photos[idx];
  count.textContent = (idx + 1) + " / " + photos.length;
  Array.from(thumbs.children).forEach((t, j) => t.classList.toggle("active", j === idx));
  const a = thumbs.children[idx];
  if (a) thumbs.scrollTo({left: a.offsetLeft - thumbs.clientWidth / 2 + a.clientWidth / 2, behavior: "smooth"});
}
function swipe(el, onLeft, onRight) {
  let sx = 0, sy = 0, multi = false;
  el.addEventListener("touchstart", e => { if (e.touches.length > 1) { multi = true; return; } multi = false; sx = e.touches[0].clientX; sy = e.touches[0].clientY; }, {passive: true});
  el.addEventListener("touchend", e => {
    if (multi) return;
    const t = e.changedTouches[0], dx = t.clientX - sx, dy = t.clientY - sy;
    if (Math.abs(dx) > 45 && Math.abs(dx) > Math.abs(dy) * 1.4) { dx < 0 ? onLeft() : onRight(); }
  }, {passive: true});
}
$("prev").onclick = () => show(idx - 1);
$("next").onclick = () => show(idx + 1);
swipe($("main"), () => show(idx + 1), () => show(idx - 1));
mainImg.onclick = openLightbox;

function openLightbox() {
  let doc;
  try { doc = window.parent.document; void doc.body; } catch (e) { window.open(photos[idx], "_blank"); return; }
  const old = doc.getElementById("ofis-lightbox"); if (old) old.remove();
  const previousOverflow = doc.body.style.overflow;
  doc.body.style.overflow = "hidden";
  const ov = doc.createElement("div");
  ov.id = "ofis-lightbox";
  ov.style.cssText = "position:fixed;top:0;left:0;right:0;bottom:0;background:rgba(0,0,0,.96);z-index:2147483000;display:flex;align-items:center;justify-content:center;";
  const img = doc.createElement("img");
  img.style.cssText = "max-width:100%;max-height:100%;object-fit:contain;user-select:none;-webkit-user-drag:none;";
  const btn = (label, css) => { const b = doc.createElement("button"); b.type = "button"; b.textContent = label; b.style.cssText = "position:absolute;border:0;border-radius:50%;width:48px;height:48px;font-size:28px;line-height:1;background:rgba(255,255,255,.9);color:#101828;cursor:pointer;padding:0;" + css; return b; };
  const closeBtn = btn("✕", "top:14px;right:14px;font-size:22px;");
  const prevBtn = btn("‹", "left:10px;top:50%;margin-top:-24px;");
  const nextBtn = btn("›", "right:10px;top:50%;margin-top:-24px;");
  const cnt = doc.createElement("div");
  cnt.style.cssText = "position:absolute;top:22px;left:50%;transform:translateX(-50%);color:#fff;font:14px system-ui,sans-serif;background:rgba(0,0,0,.55);padding:4px 12px;border-radius:12px;";
  let cur = idx;
  const render = () => { img.src = photos[cur]; cnt.textContent = (cur + 1) + " / " + photos.length; };
  const go = d => { cur = (cur + d + photos.length) % photos.length; render(); };
  const onKey = e => { if (e.key === "Escape") close(); else if (e.key === "ArrowLeft") go(-1); else if (e.key === "ArrowRight") go(1); };
  function close() {
    doc.removeEventListener("keydown", onKey);
    ov.remove();
    doc.body.style.overflow = previousOverflow;
    try { show(cur); } catch (e) {}
  }
  closeBtn.onclick = close;
  prevBtn.onclick = e => { e.stopPropagation(); go(-1); };
  nextBtn.onclick = e => { e.stopPropagation(); go(1); };
  ov.addEventListener("click", e => { if (e.target === ov) close(); });
  doc.addEventListener("keydown", onKey);
  swipe(ov, () => go(1), () => go(-1));
  ov.appendChild(img); ov.appendChild(cnt); ov.appendChild(closeBtn);
  if (photos.length > 1) { ov.appendChild(prevBtn); ov.appendChild(nextBtn); }
  doc.body.appendChild(ov);
  render();
}
show(0);
</script></body></html>"""


def render_photo_gallery(photos: list[dict]) -> None:
    urls = json.dumps([p["url"] for p in photos]).replace("</", "<\\/")
    embed_html(GALLERY_HTML.replace("__PHOTOS__", urls), 390)


def can_manage_property(row: pd.Series, user_id: str, is_admin: bool) -> bool:
    return is_admin or str(row.get("owner_id")) == user_id or str(row.get("created_by")) == user_id


def show_property_detail(client: Client, property_row: pd.Series, user_id: str, is_admin: bool, staff: pd.DataFrame, location_data: tuple[list[str], dict[str, list[str]]], favorites: set[str]) -> None:
    property_id = str(property_row.id)
    manage = can_manage_property(property_row, user_id, is_admin)
    back_col, fav_col, title_col = st.columns([1, 1.8, 4.5], vertical_alignment="center")
    if back_col.button("← Liste", key="detail_back", use_container_width=True):
        st.session_state.pop("selected_property_id", None)
        st.rerun()
    is_favorite = property_id in favorites
    fav_col.button("★ Favorilerden kaldır" if is_favorite else "☆ Favorilere ekle", key=f"fav_{property_id}", use_container_width=True,
                   type="primary" if is_favorite else "secondary", on_click=set_favorite, args=(client, user_id, property_id, not is_favorite))
    title_col.markdown(
        f"<div style='font-size:1.25rem;font-weight:700;line-height:1.25'>{escape(str(property_row.title))}</div>"
        f"<div style='color:#667085;font-size:.85rem'>{escape(str(property_row.get('property_category') or 'Daire'))} · {escape(str(property_row.get('city', '')))} / {escape(str(property_row.get('district', '')))} · İlan No: {property_id[:8].upper()}</div>",
        unsafe_allow_html=True,
    )
    left, right = st.columns([1.5, 1], gap="medium")
    try:
        photos = get_property_photos(client, property_id)
    except Exception as exc:
        photos = []
        st.warning("Fotoğraf galerisi yüklenemedi. Storage kurulumu kontrol edilmeli.")
        st.caption(str(exc))
    with left:
        if photos:
            render_photo_gallery(photos)
        else:
            st.image(str(DEFAULT_PROPERTY_IMAGE), width="stretch", alt="İlan için varsayılan fotoğraf")
            st.caption("Bu ilana henüz fotoğraf eklenmemiş.")
        tiles = info_tiles_html(property_row)
        if tiles:
            st.markdown(tiles, unsafe_allow_html=True)
    with right:
        price_value = int(float(property_row.get("price") or 0))
        status_text = _clean(property_row.get("status")) or "Aktif"
        status_color = {"Aktif": "#16a34a", "Opsiyonlu": "#d97706", "Kiralandı/Satıldı": "#64748b", "Pasif": "#dc2626"}.get(status_text, "#475569")
        chips = [_clean(property_row.get("deal_type")), _clean(property_row.get("rooms")), (_clean(property_row.get("area_m2")) + " m²") if _clean(property_row.get("area_m2")) else "", _clean(property_row.get("property_category"))]
        chip_html = "".join(f"<span class='dc-chip'>{escape(chip)}</span>" for chip in chips if chip)
        price_html = f"<div class='dc-price'>{money_text(price_value)} ₺</div>" if price_value else "<div class='dc-price muted'>Fiyat belirtilmemiş</div>"
        st.markdown(
            f"<div class='deal-card'><div class='dc-top'>{price_html}<span class='dc-status' style='background:{status_color}'>{escape(status_text)}</span></div>"
            f"<div class='dc-chips'>{chip_html}</div>"
            f"<div class='dc-row'>📍 {escape(_clean(property_row.get('city')))} / {escape(_clean(property_row.get('district')))}</div>"
            f"<div class='dc-row'>👤 Sorumlu: <b>{escape(_clean(property_row.get('sorumlu')) or 'Atanmamış')}</b></div></div>",
            unsafe_allow_html=True,
        )
        owner_id = str(property_row.get("owner_id"))
        agent = staff[staff.id.astype(str) == owner_id]
        if not agent.empty:
            person = agent.iloc[0]
            avatar = avatar_bytes(client, owner_id)
            if avatar:
                image_html = f"<img class='agent-img' src='data:image/jpeg;base64,{base64.b64encode(avatar).decode('ascii')}'>"
            else:
                image_html = f"<div class='agent-img ph'>{escape(str(person['name'])[:1].upper())}</div>"
            with st.container(key="agent_card"):
                st.markdown(f"<div class='agent-card'>{image_html}<div class='agent-name'>{escape(str(person['name']))}</div><div class='agent-role'>{escape(_clean(person.get('job_title')) or 'Gayrimenkul Danışmanı')}</div></div>", unsafe_allow_html=True)
                st.button("Danışman sayfasını aç", key="agent_open", on_click=open_employee, args=(owner_id,))
            digits = "".join(ch for ch in _clean(person.get("phone")) if ch.isdigit() or ch == "+")
            if digits:
                with st.container(key="agent_call"):
                    st.link_button("📞 Hemen Ara / İletişime Geç", f"tel:{digits}", use_container_width=True)
        detail_rows = [
            ("Brüt / net m²", f"{property_row.get('area_m2') or '—'} / {property_row.get('net_area_m2') or '—'}"),
            ("Bina yaşı", property_row.get("building_age")),
            ("Bulunduğu kat", property_row.get("floor")),
            ("Kat sayısı", property_row.get("total_floors")),
            ("Isıtma", property_row.get("heating")),
            ("Banyo", property_row.get("bathroom_count")),
            ("Mutfak", property_row.get("kitchen")),
            ("Balkon", "Var" if property_row.get("balcony") else "Yok"),
            ("Asansör", "Var" if property_row.get("elevator") else "Yok"),
            ("Otopark", property_row.get("parking")),
            ("Eşyalı", "Evet" if property_row.get("furnished") else "Hayır"),
            ("Kullanım durumu", property_row.get("usage_status")),
            ("Site içerisinde", "Evet" if property_row.get("within_complex") else "Hayır"),
            ("Site adı", property_row.get("complex_name")),
            ("Aidat", f"{int(property_row.get('monthly_dues') or 0):,}".replace(",", ".") + " ₺"),
            ("Tapu durumu", property_row.get("title_deed_status")),
            ("Krediye uygun", "Evet" if property_row.get("loan_eligible") else "Hayır"),
            ("Takas", "Evet" if property_row.get("swap_allowed") else "Hayır"),
        ]
        detail_rows = [(label, value) for label, value in detail_rows if value not in (None, "") and not (isinstance(value, float) and pd.isna(value))]
        table_rows = "".join(f"<tr><th>{escape(str(label))}</th><td>{escape(str(value))}</td></tr>" for label, value in detail_rows)
        st.markdown(f"<table class='property-feature-table'>{table_rows}</table>", unsafe_allow_html=True)

    tab_icons = {"Açıklama": "📝", "Konum": "📍", "Fotoğraflar": "🖼️", "İlan teklifleri": "💰", "Mesaj": "✉️", "İndir ve paylaş": "📤", "Düzenle": "✏️", "Sil": "🗑️"}
    tab_names = ["Açıklama", "Konum", "Fotoğraflar", "İlan teklifleri", "Mesaj", "İndir ve paylaş"] + (["Düzenle", "Sil"] if manage else [])
    detail_tabs = dict(zip(tab_names, st.tabs([f"{tab_icons[name]} {name}" for name in tab_names])))
    with detail_tabs["Açıklama"]:
        st.write(property_row.get("description") or property_row.get("notes") or "Bu ilan için açıklama girilmemiş.")
        with st.expander("📈 Fiyat geçmişi"):
            price_history_table(client, property_id)
    with detail_tabs["Konum"]:
        st.write(f"{property_row.get('city', '')} / {property_row.get('district', '')} — {property_row.get('address') or 'Açık adres belirtilmemiş.'}")
        latitude, longitude = property_row.get("latitude"), property_row.get("longitude")
        has_coordinates = (
            pd.notna(latitude) and pd.notna(longitude)
            and str(latitude).strip() and str(longitude).strip()
        )
        map_query = f"{latitude},{longitude}" if has_coordinates else ", ".join(
            str(value) for value in (property_row.get("address"), property_row.get("district"), property_row.get("city"))
            if pd.notna(value) and str(value).strip()
        )
        if map_query:
            embed_url = f"https://maps.google.com/maps?q={quote(map_query)}&output=embed"
            embed_page(embed_url, 300)
        maps_url = property_row.get("maps_url")
        if pd.notna(maps_url) and str(maps_url).strip():
            st.link_button("Google Haritalar'da aç", str(maps_url), use_container_width=True)
        elif has_coordinates:
            map_url = f"https://www.google.com/maps/search/?api=1&query={latitude}%2C{longitude}"
            st.link_button("Google Haritalar'da aç", map_url, use_container_width=True)
        distance_tool(property_row)
    with detail_tabs["Fotoğraflar"]:
        photos_tab_content(client, property_id, photos, manage)
    with detail_tabs["İlan teklifleri"]:
        offers_panel(client, property_id, user_id, is_admin, staff, str(property_row.get("owner_id")), "detail")
    with detail_tabs["Mesaj"]:
        owner_id = str(property_row.get("owner_id"))
        if owner_id == user_id:
            st.info("Bu ilanın sorumlusu sizsiniz. İlanla ilgili mesajlar Mesajlar sekmesine gelir.")
        else:
            compose_message_form(client, user_id, staff, f"detail_{property_id}", default_recipient=owner_id, lock_recipient=True, property_id=property_id, subject_hint=f"İlan: {_clean(property_row.get('title'))}")
    with detail_tabs["İndir ve paylaş"]:
        download_tab_content(property_row, photos)
    if manage:
        with detail_tabs["Düzenle"]:
            edit_tab_content(client, property_row, is_admin, staff, location_data)
        with detail_tabs["Sil"]:
            delete_tab_content(client, property_row)
    st.markdown("<div style='height:2.5rem'></div>", unsafe_allow_html=True)


USER_AGENT = {"User-Agent": "ofis-portfoy-streamlit/1.0"}
TURKEY_CENTER = (39.0, 35.0)


def _norm(text) -> str:
    text = str(text or "").replace("İ", "i").replace("I", "ı").lower()
    text = "".join(ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", " ", text.replace("ı", "i")).strip()


@st.cache_data(ttl=3600, show_spinner=False)
def reverse_geocode(lat: float, lng: float) -> dict:
    response = requests.get(
        "https://nominatim.openstreetmap.org/reverse",
        params={"lat": lat, "lon": lng, "format": "jsonv2", "addressdetails": 1, "zoom": 18, "accept-language": "tr"},
        headers=USER_AGENT, timeout=10,
    )
    response.raise_for_status()
    return response.json()


@st.cache_data(ttl=3600, show_spinner=False)
def search_location(query: str) -> dict | None:
    response = requests.get(
        "https://nominatim.openstreetmap.org/search",
        params={"q": query, "format": "jsonv2", "addressdetails": 1, "limit": 1, "countrycodes": "tr", "accept-language": "tr"},
        headers=USER_AGENT, timeout=10,
    )
    response.raise_for_status()
    data = response.json()
    return data[0] if data else None


def describe_location(result: dict, provinces: list[str], districts_by_province: dict[str, list[str]]) -> dict:
    addr = result.get("address") or {}
    province = None
    province_lookup = {_norm(p): p for p in provinces}
    for field in ("province", "state", "city", "region", "county"):
        value = addr.get(field)
        if value:
            province = province_lookup.get(_norm(re.sub(r"\s+(ili|province)$", "", str(value), flags=re.I)))
            if province:
                break
    district = None
    if province:
        district_lookup = {_norm(d): d for d in districts_by_province.get(province, [])}
        for field in ("county", "city_district", "district", "town", "municipality", "suburb", "city", "village"):
            value = addr.get(field)
            if value:
                cleaned = re.sub(r"\s+(belediyesi|ilçesi|merkez)$", "", str(value), flags=re.I)
                district = district_lookup.get(_norm(cleaned)) or district_lookup.get(_norm(value))
                if district:
                    break
    neighbourhood = addr.get("neighbourhood") or addr.get("quarter") or addr.get("suburb")
    street = addr.get("road") or addr.get("pedestrian") or addr.get("footway")
    if street and addr.get("house_number"):
        street = f"{street} No: {addr['house_number']}"
    address = ", ".join(part for part in (neighbourhood, street) if part)
    if not address:
        address = ", ".join(part.strip() for part in str(result.get("display_name", "")).split(",")[:3])
    return {"city": province, "district": district, "address": address}


def clear_form_state(prefix: str) -> None:
    for state_key in list(st.session_state.keys()):
        if state_key.startswith(prefix + "_") and not state_key.endswith("_last_click"):
            try:
                del st.session_state[state_key]
            except Exception:
                pass


def property_form(key: str, values: dict | None, location_data: tuple[list[str], dict[str, list[str]]], staff: pd.DataFrame, default_owner: str | None = None, owner_editable: bool = True, extra_tab: tuple | None = None) -> dict:
    provinces, districts_by_province = location_data
    values = values or {}
    is_new = not values

    def K(name: str) -> str:
        return f"{key}_{name}"

    def raw(name: str, default=None):
        value = values.get(name)
        if value is None:
            return default
        try:
            if pd.isna(value):
                return default
        except (TypeError, ValueError):
            pass
        return value

    def init(name: str, value) -> None:
        if K(name) not in st.session_state:
            st.session_state[K(name)] = value

    def pick(col, label: str, name: str, options: list[str], default: str, db_name: str | None = None):
        current = str(raw(db_name or name, default))
        opts = list(options)
        if current not in opts:
            opts.append(current)
        init(name, current)
        return col.selectbox(label, opts, key=K(name))

    def number(col, label: str, name: str, step: int = 1, default: int = 0, db_name: str | None = None):
        init(name, int(raw(db_name or name, default) or 0))
        return col.number_input(label, min_value=0, step=step, key=K(name))

    def check(col, label: str, name: str):
        init(name, bool(raw(name, False)))
        return col.checkbox(label, key=K(name))

    tab_names = ["Temel bilgiler", "Konum ve adres", "Özellikler", "Açıklama"] + ([extra_tab[0]] if extra_tab else [])
    tabs = st.tabs(tab_names)
    tab_basic, tab_loc, tab_feat, tab_desc = tabs[:4]

    with tab_basic:
        c1, c2, c3 = st.columns(3)
        init("title", str(raw("title", "")))
        title = c1.text_input("İlan başlığı *", key=K("title"), placeholder="Deniz manzaralı 2+1 daire")
        category = pick(c2, "Emlak tipi", "category", ["Daire", "Rezidans", "Müstakil ev", "Villa", "Yazlık", "İşyeri", "Arsa", "Bina", "Turistik tesis"], "Daire", "property_category")
        deal_type = pick(c3, "İlan tipi", "deal_type", ["Kiralık", "Satılık"], "Kiralık")
        c4, c5, c6 = st.columns(3)
        price = number(c4, "Fiyat (₺)", "price", 1000)
        rooms = pick(c5, "Oda sayısı", "rooms", ["Belirtilmedi", "Stüdyo (1+0)", "1+1", "2+1", "3+1", "4+1", "5+1", "Diğer"], "Belirtilmedi")
        status = pick(c6, "İlan durumu", "status", ["Aktif", "Opsiyonlu", "Kiralandı/Satıldı", "Pasif"], "Aktif")
        c7, c8, c9 = st.columns(3)
        gross_area = number(c7, "Brüt m²", "area_m2", 5)
        net_area = number(c8, "Net m²", "net_area_m2", 5)
        building_age = pick(c9, "Bina yaşı", "building_age", ["Belirtilmemiş", "0", "1", "2", "3", "4", "5-10", "11-15", "16-20", "21+"], "Belirtilmemiş")

    with tab_loc:
        st.caption("Haritada ilanın yerine dokunarak işaretleyin; il, ilçe ve adres otomatik dolar. İsterseniz önce adres arayabilirsiniz.")
        init("lat", float(raw("latitude")) if raw("latitude") is not None else None)
        init("lng", float(raw("longitude")) if raw("longitude") is not None else None)

        def apply_pin(new_lat: float, new_lng: float, geo: dict | None = None) -> bool:
            st.session_state[K("lat")] = float(new_lat)
            st.session_state[K("lng")] = float(new_lng)
            try:
                geo = geo or reverse_geocode(round(float(new_lat), 6), round(float(new_lng), 6))
                info = describe_location(geo, provinces, districts_by_province)
            except Exception:
                return False
            if info["address"]:
                st.session_state[K("address")] = info["address"]
            if info["city"] and provinces:
                st.session_state[K("city")] = info["city"]
                if info["district"]:
                    st.session_state[K(f"district_{info['city']}")] = info["district"]
            return True

        search_col, button_col = st.columns([4, 1], vertical_alignment="bottom")
        query = search_col.text_input("Adres veya yer ara (isteğe bağlı)", key=K("search"), placeholder="Örn. Yeni Meram, Konya")
        if button_col.button("Haritada bul", key=K("search_btn"), use_container_width=True):
            if not query.strip():
                st.warning("Aranacak adresi yazın.")
            else:
                try:
                    found = search_location(query.strip() + ", Türkiye")
                except Exception:
                    found = None
                    st.warning("Adres arama servisine şu anda ulaşılamadı. Haritaya dokunarak işaretleyebilirsiniz.")
                if found:
                    apply_pin(float(found["lat"]), float(found["lon"]), geo=found)
                    st.rerun()
                elif found is None:
                    st.warning("Adres bulunamadı. Haritada doğrudan dokunarak işaretleyebilirsiniz.")

        lat, lng = st.session_state.get(K("lat")), st.session_state.get(K("lng"))
        has_pin = lat is not None and lng is not None
        if has_pin:
            center, zoom = (lat, lng), 16
        else:
            center, zoom = TURKEY_CENTER, 6
            city_now = st.session_state.get(K("city")) or raw("city") or ("Konya" if "Konya" in provinces else None)
            district_now = st.session_state.get(K(f"district_{city_now}")) or raw("district")
            if city_now:
                try:
                    found = search_location(", ".join(part for part in (district_now, city_now, "Türkiye") if part))
                    if found:
                        center, zoom = (float(found["lat"]), float(found["lon"])), (12 if district_now else 10)
                except Exception:
                    pass
        fmap = folium.Map(location=center, zoom_start=zoom, control_scale=True)
        if has_pin:
            folium.Marker([lat, lng], tooltip="Seçilen konum").add_to(fmap)
        out = st_folium(fmap, height=300, use_container_width=True, key=K("map"), returned_objects=["last_clicked"])
        clicked = (out or {}).get("last_clicked")
        if clicked and clicked != st.session_state.get(K("last_click")):
            st.session_state[K("last_click")] = clicked
            if not apply_pin(clicked["lat"], clicked["lng"]):
                st.session_state[K("geo_failed")] = True
            else:
                st.session_state.pop(K("geo_failed"), None)
            st.rerun()
        if has_pin:
            info_col, clear_col = st.columns([4, 1], vertical_alignment="center")
            info_col.caption(f"📌 İşaretlenen konum: {lat:.5f}, {lng:.5f}")
            if clear_col.button("Konumu kaldır", key=K("clear_pin"), use_container_width=True):
                st.session_state[K("lat")] = None
                st.session_state[K("lng")] = None
                st.rerun()
            if st.session_state.get(K("geo_failed")):
                st.warning("Bu noktanın adresi otomatik alınamadı; il, ilçe ve adresi elle girin. Koordinat yine de kaydedilecek.")

        c10, c11, c12 = st.columns([1, 1, 2])
        default_city = "Konya" if "Konya" in provinces else (provinces[0] if provinces else "")
        if provinces:
            province_lookup = {_norm(p): p for p in provinces}
            init("city", province_lookup.get(_norm(raw("city", "")), default_city))
            city = c10.selectbox("İl *", provinces, key=K("city"))
            district_options = districts_by_province.get(city, [])
        else:
            init("city", str(raw("city", "")))
            city = c10.text_input("İl *", key=K("city"))
            district_options = []
        if district_options:
            district_lookup = {_norm(d): d for d in district_options}
            guess = district_lookup.get(_norm(raw("district", ""))) or ("Meram" if is_new and city == "Konya" and "Meram" in district_options else district_options[0])
            init(f"district_{city}", guess)
            district = c11.selectbox("İlçe *", district_options, key=K(f"district_{city}"))
        else:
            init(f"district_{city}", str(raw("district", "")))
            district = c11.text_input("İlçe *", key=K(f"district_{city}"))
        init("address", str(raw("address", "")))
        address = c12.text_input("Açık adres / konum tarifi", key=K("address"), placeholder="Mahalle, cadde, bina no (haritadan otomatik dolar)")

    with tab_feat:
        f1, f2, f3, f4 = st.columns(4)
        floor = pick(f1, "Bulunduğu kat", "floor", ["Belirtilmemiş", "Bodrum", "Giriş", "Bahçe katı", "Yüksek giriş", "1", "2", "3", "4", "5", "6-10", "11+", "Çatı"], "Belirtilmemiş")
        total_floors = number(f2, "Binadaki kat sayısı", "total_floors")
        heating = pick(f3, "Isıtma", "heating", ["Belirtilmemiş", "Kombi (Doğalgaz)", "Merkezi", "Yerden ısıtma", "Klima", "Soba", "Diğer"], "Belirtilmemiş")
        bathrooms = number(f4, "Banyo sayısı", "bathroom_count", 1, default=1)
        g1, g2, g3, g4 = st.columns(4)
        kitchen = pick(g1, "Mutfak", "kitchen", ["Belirtilmemiş", "Açık (Amerikan)", "Kapalı"], "Belirtilmemiş")
        parking = pick(g2, "Otopark", "parking", ["Yok", "Açık", "Kapalı", "Açık ve kapalı"], "Yok")
        usage_status = pick(g3, "Kullanım durumu", "usage_status", ["Boş", "Kiracılı", "Mal sahibi oturuyor"], "Boş")
        title_deed = pick(g4, "Tapu durumu", "title_deed_status", ["Belirtilmemiş", "Kat mülkiyetli", "Kat irtifaklı", "Arsa tapulu", "Hisseli", "Diğer"], "Belirtilmemiş")
        h1, h2, h3, h4 = st.columns(4, vertical_alignment="bottom")
        init("complex_name", str(raw("complex_name", "")))
        complex_name = h1.text_input("Site adı", key=K("complex_name"))
        dues = number(h2, "Aylık aidat (₺)", "monthly_dues", 100)
        within_complex = check(h3, "Site içerisinde", "within_complex")
        furnished = check(h4, "Eşyalı", "furnished")
        i1, i2, i3, i4 = st.columns(4)
        balcony = check(i1, "Balkon", "balcony")
        elevator = check(i2, "Asansör", "elevator")
        loan_eligible = check(i3, "Krediye uygun", "loan_eligible")
        swap_allowed = check(i4, "Takas değerlendirilir", "swap_allowed")

    with tab_desc:
        init("description", str(raw("description", "")))
        description = st.text_area("İlan açıklaması", height=130, key=K("description"))
        owner_id = None
        if owner_editable and not staff.empty:
            owner_ids = staff["id"].astype(str).tolist()
            stored_owner = str(raw("owner_id", default_owner or owner_ids[0]))
            init("owner", stored_owner if stored_owner in owner_ids else owner_ids[0])
            owner_id = st.selectbox("Sorumlu çalışan", owner_ids, format_func=lambda x: staff.loc[staff.id.astype(str) == x, "name"].iloc[0], key=K("owner"))

    if extra_tab:
        with tabs[4]:
            extra_tab[1]()

    payload = {
        "title": title.strip(), "city": city, "district": district, "deal_type": deal_type,
        "price": int(price), "rooms": rooms, "area_m2": int(gross_area), "net_area_m2": int(net_area),
        "property_category": category, "building_age": building_age, "floor": floor,
        "total_floors": int(total_floors), "heating": heating, "bathroom_count": int(bathrooms),
        "kitchen": kitchen, "balcony": balcony, "elevator": elevator, "parking": parking,
        "furnished": furnished, "usage_status": usage_status, "within_complex": within_complex,
        "complex_name": complex_name, "monthly_dues": int(dues), "title_deed_status": title_deed,
        "loan_eligible": loan_eligible, "swap_allowed": swap_allowed, "address": address,
        "description": description, "status": status,
    }
    if has_pin:
        payload["maps_url"] = f"https://www.google.com/maps/search/?api=1&query={lat}%2C{lng}"
        payload["latitude"], payload["longitude"] = float(lat), float(lng)
    else:
        google_query = ", ".join(part for part in [address, district, city, "Türkiye"] if part)
        payload["maps_url"] = f"https://www.google.com/maps/search/?api=1&query={quote(google_query)}"
        payload["latitude"], payload["longitude"] = None, None
    if owner_id:
        payload["owner_id"] = owner_id
    return payload


def delete_properties(db_client: Client, ids: list[str]) -> tuple[int, list[str]]:
    deleted, errors = 0, []
    for start in range(0, len(ids), 50):
        chunk = ids[start:start + 50]
        try:
            result = db_client.table("properties").delete().in_("id", chunk).execute()
        except Exception as exc:
            errors.append(f"İlanlar silinemedi: {exc}")
            continue
        removed = [str(r["id"]) for r in (result.data or [])]
        deleted += len(removed)
        for property_id in removed:
            try:
                bucket = service_client(db_client).storage.from_(PHOTO_BUCKET)
                paths = [f"{property_id}/{o['name']}" for o in (bucket.list(property_id, {"limit": 1000}) or []) if o.get("id")]
                if paths:
                    bucket.remove(paths)
            except Exception as exc:
                errors.append(f"{property_id[:8].upper()} numaralı ilanın fotoğrafları silinemedi: {exc}")
    return deleted, errors


def add_property(client: Client, user_id: str, staff: pd.DataFrame, location_data: tuple[list[str], dict[str, list[str]]]) -> None:
    holder: dict = {}

    def photo_tab() -> None:
        holder["files"] = st.file_uploader("İlan fotoğrafları (isteğe bağlı · JPG, PNG veya WebP; her biri en fazla 10 MB, en fazla 20 adet)", type=["jpg", "jpeg", "png", "webp"], accept_multiple_files=True, key="np_photos")
        st.caption("Fotoğraf eklemezseniz ilanda varsayılan görsel gösterilir; fotoğrafı sonradan ilan detayından ekleyebilirsiniz.")

    with st.container(border=True):
        payload = property_form("np", {}, location_data, staff, default_owner=user_id, extra_tab=("Fotoğraflar", photo_tab))
        photos = holder.get("files")
        saved = st.button("İlanı kaydet", type="primary", use_container_width=True, key="save_new_property")
    if saved:
        if not payload["title"] or not payload["city"] or not payload["district"]:
            st.error("Başlık (Temel bilgiler), il ve ilçe (Konum ve adres) alanlarını doldurun.")
            return
        if photos and len(photos) > 20:
            st.error("Bir ilana en fazla 20 fotoğraf ekleyebilirsiniz.")
            return
        invalid = [f.name for f in (photos or []) if f.type not in {"image/jpeg", "image/png", "image/webp"} or len(f.getvalue()) > MAX_PHOTO_BYTES]
        if invalid:
            st.error("Fotoğrafları kontrol edin (JPG, PNG veya WebP; her biri en fazla 10 MB): " + ", ".join(invalid))
            return
        payload["created_by"] = user_id
        try:
            result = client.table("properties").insert(payload).select("id").single().execute()
            property_id = str(result.data["id"])
            count, errors = upload_property_photos(client, property_id, photos) if photos else (0, [])
            st.session_state.property_save_notice = {"uploaded": count, "errors": errors}
            st.session_state.selected_property_id = property_id
            clear_form_state("np")
            go_to_section("Daireler")
            st.rerun()
        except Exception as exc:
            st.error("İlan kaydedilemedi. Yeni ilan alanlarının Supabase migration'ının çalıştırıldığını kontrol edin.")
            st.caption(str(exc))


def edit_tab_content(client: Client, row: pd.Series, is_admin: bool, staff: pd.DataFrame, location_data: tuple[list[str], dict[str, list[str]]]) -> None:
    property_id = str(row.id)
    prefix = f"edit_{property_id}"
    payload = property_form(prefix, row.to_dict(), location_data, staff, owner_editable=is_admin)
    if st.button("Değişiklikleri kaydet", type="primary", use_container_width=True, key=f"{prefix}_save"):
        if not payload["title"] or not payload["city"] or not payload["district"]:
            st.error("Başlık, il ve ilçe alanlarını doldurun.")
        else:
            try:
                client.table("properties").update(payload).eq("id", property_id).execute()
                clear_form_state(prefix)
                st.success("İlan güncellendi.")
                st.rerun()
            except Exception as exc:
                st.error("İlan güncellenemedi.")
                st.caption(str(exc))


def delete_tab_content(client: Client, row: pd.Series) -> None:
    property_id = str(row.id)
    prefix = f"edit_{property_id}"
    st.warning("İlan ve fotoğrafları kalıcı olarak silinir. Bu işlem geri alınamaz.")
    confirm = st.checkbox("Bu ilanı silmek istediğimi onaylıyorum.", key=f"{prefix}_confirm_delete")
    if st.button("İlanı kalıcı olarak sil", type="primary", disabled=not confirm, key=f"{prefix}_delete"):
        deleted, errors = delete_properties(service_client(client), [property_id])
        if deleted:
            clear_form_state(prefix)
            st.session_state.pop("selected_property_id", None)
            st.session_state["property_delete_notice"] = "İlan silindi."
            st.rerun()
        elif errors:
            for error in errors:
                st.error(error)
        else:
            st.error("İlan silinemedi. Supabase'de bu tablo için silme (DELETE) izni tanımlı olmayabilir; yöneticinizle iletişime geçin.")


def photos_tab_content(client: Client, property_id: str, photos: list[dict], manage: bool) -> None:
    more_photos = st.file_uploader("Bu ilana fotoğraf ekle", type=["jpg", "jpeg", "png", "webp"], accept_multiple_files=True, key=f"detail_upload_{property_id}")
    if st.button("Fotoğrafları yükle", disabled=not more_photos, key=f"detail_upload_button_{property_id}"):
        count, errors = upload_property_photos(client, property_id, more_photos or [])
        for error in errors:
            st.error(error)
        if count:
            thumb_store().pop(property_id, None)
            st.success(f"{count} fotoğraf yüklendi.")
            st.rerun()
    if not manage:
        return
    st.divider()
    if not photos:
        st.caption("Bu ilanda silinebilecek fotoğraf yok.")
        return
    generation = st.session_state.get(f"photo_generation_{property_id}", 0)
    st.markdown("**Fotoğrafları sil** — silmek istediklerinizi işaretleyin")
    chosen: list[str] = []
    with st.container(height=280, key=f"photo_grid_{property_id}"):
        columns = st.columns(4)
        for index, photo in enumerate(photos):
            with columns[index % 4]:
                st.image(photo["url"], width=120)
                if st.checkbox(f"Sil · {index + 1}", key=f"photo_del_{property_id}_{generation}_{index}"):
                    chosen.append(photo["name"])
    if st.button(f"Seçili {len(chosen)} fotoğrafı sil", type="primary", disabled=not chosen, key=f"photo_delete_button_{property_id}_{generation}"):
        try:
            removed = service_client(client).storage.from_(PHOTO_BUCKET).remove([f"{property_id}/{name}" for name in chosen])
        except Exception as exc:
            st.error("Fotoğraflar silinemedi.")
            st.caption(str(exc))
            return
        if not removed:
            st.error("Fotoğraflar silinemedi. Streamlit Secrets'a service_role_key ekleyin veya Supabase Storage'da DELETE izni tanımlayın.")
            return
        thumb_store().pop(property_id, None)
        st.session_state[f"photo_generation_{property_id}"] = generation + 1
        st.session_state["property_delete_notice"] = f"{len(removed)} fotoğraf silindi."
        st.rerun()


def _clean(value) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def property_fact_rows(row: pd.Series) -> list[tuple[str, str]]:
    def money(name: str) -> str:
        value = row.get(name)
        try:
            return f"{int(value):,}".replace(",", ".") + " ₺" if value not in (None, "") and not pd.isna(value) else ""
        except (TypeError, ValueError):
            return ""

    def yes_no(name: str, yes: str = "Evet", no: str = "Hayır") -> str:
        return yes if bool(row.get(name)) else no

    facts = [
        ("İlan no", str(row.id)[:8].upper()), ("Başlık", _clean(row.get("title"))), ("İşlem", _clean(row.get("deal_type"))),
        ("Emlak tipi", _clean(row.get("property_category"))), ("Fiyat", money("price")), ("İl", _clean(row.get("city"))),
        ("İlçe", _clean(row.get("district"))), ("Adres", _clean(row.get("address"))), ("Oda sayısı", _clean(row.get("rooms"))),
        ("Brüt m²", _clean(row.get("area_m2"))), ("Net m²", _clean(row.get("net_area_m2"))), ("Bina yaşı", _clean(row.get("building_age"))),
        ("Bulunduğu kat", _clean(row.get("floor"))), ("Binadaki kat sayısı", _clean(row.get("total_floors"))), ("Isıtma", _clean(row.get("heating"))),
        ("Banyo sayısı", _clean(row.get("bathroom_count"))), ("Mutfak", _clean(row.get("kitchen"))), ("Balkon", yes_no("balcony", "Var", "Yok")),
        ("Asansör", yes_no("elevator", "Var", "Yok")), ("Otopark", _clean(row.get("parking"))), ("Eşyalı", yes_no("furnished")),
        ("Kullanım durumu", _clean(row.get("usage_status"))), ("Site içerisinde", yes_no("within_complex")), ("Site adı", _clean(row.get("complex_name"))),
        ("Aylık aidat", money("monthly_dues")), ("Tapu durumu", _clean(row.get("title_deed_status"))), ("Krediye uygun", yes_no("loan_eligible")),
        ("Takas", yes_no("swap_allowed")), ("Durum", _clean(row.get("status"))), ("Sorumlu", _clean(row.get("sorumlu"))),
        ("Harita bağlantısı", _clean(row.get("maps_url"))), ("Açıklama", _clean(row.get("description"))),
    ]
    return [(label, value) for label, value in facts if value]


def build_excel(facts: list[tuple[str, str]], images: list | None = None, office_name: str = DEFAULT_OFFICE_NAME) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from PIL import Image, ImageOps

    images = images or []
    data = dict(facts)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "İlan"
    sheet.sheet_view.showGridLines = False
    for index in range(1, 13):
        sheet.column_dimensions[get_column_letter(index)].width = 8.9
    face = "Arial"
    navy, blue, gray, ink = "1E3A8A", "1D4ED8", "667085", "101828"

    def block(first_col: int, last_col: int, row: int, value: str, *, size=11, bold=False, color=ink, fill=None, horizontal="left", vertical="center", wrap=False, shrink=False) -> None:
        start, end = f"{get_column_letter(first_col)}{row}", f"{get_column_letter(last_col)}{row}"
        if first_col != last_col:
            sheet.merge_cells(f"{start}:{end}")
        cell = sheet[start]
        cell.value = value
        cell.font = Font(name=face, size=size, bold=bold, color=color)
        cell.alignment = Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap, shrink_to_fit=shrink)
        if fill:
            for col in range(first_col, last_col + 1):
                sheet.cell(row=row, column=col).fill = PatternFill("solid", fgColor=fill)

    row = 1
    sheet.row_dimensions[row].height = 32
    block(1, 8, row, "  " + office_name, size=16, bold=True, color="FFFFFF", fill=navy)
    block(9, 12, row, f"İlan No: {data.get('İlan no', '')}  ", size=11, color="DBEAFE", fill=navy, horizontal="right")
    row += 1
    sheet.row_dimensions[row].height = 8
    row += 1
    title = data.get("Başlık", "İlan")
    sheet.row_dimensions[row].height = 26 * max(1, -(-len(title) // 52))
    block(1, 12, row, title, size=18, bold=True, wrap=True)
    row += 1
    subtitle = " · ".join(part for part in (data.get("İşlem"), data.get("Emlak tipi"), " / ".join(p for p in (data.get("İl"), data.get("İlçe")) if p)) if part)
    block(1, 12, row, subtitle, size=11, color="475467")
    row += 1
    sheet.row_dimensions[row].height = 38
    block(1, 9, row, data.get("Fiyat", ""), size=22, bold=True, color=blue)
    if data.get("Durum"):
        block(10, 12, row, data["Durum"], size=12, bold=True, color=navy, fill="E0E7FF", horizontal="center")
    row += 1
    sheet.row_dimensions[row].height = 8
    row += 1

    if images:
        main_rows = 15
        for offset in range(main_rows):
            sheet.row_dimensions[row + offset].height = 15
        buffer = io.BytesIO()
        ImageOps.fit(images[0], (780, 300), method=Image.LANCZOS).save(buffer, "JPEG", quality=85)
        main_image = XLImage(io.BytesIO(buffer.getvalue()))
        main_image.width, main_image.height = 780, 300
        sheet.add_image(main_image, f"A{row}")
        row += main_rows + 1
        extra = images[1:4]
        if extra:
            thumb_rows = 8
            for offset in range(thumb_rows):
                sheet.row_dimensions[row + offset].height = 15
            for index, picture in enumerate(extra):
                buffer = io.BytesIO()
                ImageOps.fit(picture, (248, 160), method=Image.LANCZOS).save(buffer, "JPEG", quality=85)
                thumb = XLImage(io.BytesIO(buffer.getvalue()))
                thumb.width, thumb.height = 248, 160
                sheet.add_image(thumb, f"{get_column_letter(1 + index * 4)}{row}")
            row += thumb_rows + 1
    else:
        for offset in range(5):
            sheet.row_dimensions[row + offset].height = 20
        sheet.merge_cells(start_row=row, start_column=1, end_row=row + 4, end_column=12)
        cell = sheet.cell(row=row, column=1, value="Bu ilana fotoğraf eklenmemiş")
        cell.font = Font(name=face, size=12, color="98A2B3")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for r in range(row, row + 5):
            for c in range(1, 13):
                sheet.cell(row=r, column=c).fill = PatternFill("solid", fgColor="F2F4F7")
        row += 6

    skip = {"İlan no", "Başlık", "Açıklama", "Adres", "Harita bağlantısı", "Fiyat", "Sorumlu", "Durum", "İşlem", "İl", "İlçe", "Emlak tipi"}
    grid = [(label, value) for label, value in facts if label not in skip]
    for index in range(0, len(grid), 2):
        sheet.row_dimensions[row].height = 22
        shade = "F9FAFB" if (index // 2) % 2 == 0 else None
        for side, (first_col, item) in enumerate(((1, grid[index]), (7, grid[index + 1] if index + 1 < len(grid) else None))):
            if item is None:
                continue
            block(first_col, first_col + 1, row, "  " + item[0], size=10, color=gray, fill=shade)
            block(first_col + 2, first_col + 5, row, item[1], size=10, bold=True, fill=shade, shrink=True)
        row += 1
    row += 1

    if data.get("Adres"):
        sheet.row_dimensions[row].height = 18 * max(1, -(-len(data["Adres"]) // 95))
        block(1, 12, row, "Adres: " + data["Adres"], size=10, color="344054", wrap=True, vertical="top")
        row += 2
    if data.get("Açıklama"):
        block(1, 12, row, "Açıklama", size=12, bold=True)
        row += 1
        text = data["Açıklama"]
        lines = sum(max(1, -(-len(part) // 105)) for part in text.splitlines() or [""])
        sheet.row_dimensions[row].height = min(400, 15 * lines + 8)
        block(1, 12, row, text, size=10, color="344054", wrap=True, vertical="top")
        row += 2
    sheet.row_dimensions[row].height = 22
    block(1, 12, row, f"Sorumlu: {data.get('Sorumlu', '—')}   ·   Hazırlanma tarihi: {date.today().strftime('%d.%m.%Y')}", size=10, color=gray, horizontal="center")

    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _font(size: int, bold: bool = False, note: bool = False):
    from PIL import ImageFont

    name = "DejaVuSerif-BoldItalic.ttf" if note else ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf")
    candidates = [Path(__file__).parent / "assets" / name]
    candidates += [Path(folder) / name for folder in ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "/usr/share/fonts/TTF")]
    try:
        import matplotlib
        candidates.append(Path(matplotlib.get_data_path()) / "fonts" / "ttf" / name)
    except Exception:
        pass
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    try:
        return ImageFont.load_default(size)
    except TypeError:
        return ImageFont.load_default()


def _wrap_text(draw, text: str, font, max_width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in str(text).splitlines() or [""]:
        current = ""
        for word in paragraph.split():
            trial = f"{current} {word}".strip()
            if draw.textlength(trial, font=font) <= max_width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = word
        lines.append(current)
    return lines


def _fetch_images(photos: list[dict], limit: int = 4) -> list:
    from PIL import Image, ImageOps

    images = []
    for photo in photos[:limit]:
        try:
            response = requests.get(photo["url"], timeout=20)
            response.raise_for_status()
            image = ImageOps.exif_transpose(Image.open(io.BytesIO(response.content))).convert("RGB")
            image.thumbnail((1600, 1600))
            images.append(image)
        except Exception:
            continue
    return images


def render_listing_image(facts: list[tuple[str, str]], images: list, office_name: str = DEFAULT_OFFICE_NAME):
    from PIL import Image, ImageDraw, ImageOps

    data = dict(facts)
    width, height, margin = 1240, 1754, 60
    page = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(page)
    inner = width - 2 * margin

    draw.rectangle([0, 0, width, 110], fill="#1e3a8a")
    draw.text((margin, 34), office_name, font=_font(40, True), fill="white")
    id_text = f"İlan No: {data.get('İlan no', '')}"
    id_font = _font(30)
    draw.text((width - margin - draw.textlength(id_text, font=id_font), 44), id_text, font=id_font, fill="#dbeafe")

    y = 140
    title_font = _font(46, True)
    for line in _wrap_text(draw, data.get("Başlık", "İlan"), title_font, inner)[:2]:
        draw.text((margin, y), line, font=title_font, fill="#101828")
        y += 60
    subtitle = " · ".join(part for part in (data.get("İşlem"), data.get("Emlak tipi"), " / ".join(p for p in (data.get("İl"), data.get("İlçe")) if p)) if part)
    draw.text((margin, y + 4), subtitle, font=_font(28), fill="#475467")
    y += 54

    price_font = _font(62, True)
    draw.text((margin, y), data.get("Fiyat", ""), font=price_font, fill="#1d4ed8")
    status = data.get("Durum", "")
    if status:
        badge_font = _font(28, True)
        badge_width = int(draw.textlength(status, font=badge_font)) + 44
        draw.rounded_rectangle([width - margin - badge_width, y + 10, width - margin, y + 62], radius=26, fill="#e0e7ff")
        draw.text((width - margin - badge_width + 22, y + 18), status, font=badge_font, fill="#1e3a8a")
    y += 96

    main_height = 410
    if images:
        page.paste(ImageOps.fit(images[0], (inner, main_height), method=Image.LANCZOS), (margin, y))
        y += main_height + 12
        extra = images[1:4]
        if extra:
            gap = 12
            thumb_width = (inner - gap * (len(extra) - 1)) // len(extra)
            for index, image in enumerate(extra):
                page.paste(ImageOps.fit(image, (thumb_width, 170), method=Image.LANCZOS), (margin + index * (thumb_width + gap), y))
            y += 170 + 12
    else:
        draw.rectangle([margin, y, width - margin, y + 220], fill="#f2f4f7")
        placeholder_font = _font(32)
        text = "Bu ilana fotoğraf eklenmemiş"
        draw.text(((width - draw.textlength(text, font=placeholder_font)) / 2, y + 90), text, font=placeholder_font, fill="#98a2b3")
        y += 220 + 12
    y += 8

    skip = {"İlan no", "Başlık", "Açıklama", "Adres", "Harita bağlantısı", "Fiyat", "Sorumlu", "Durum", "İşlem", "İl", "İlçe", "Emlak tipi"}
    grid = [(label, value) for label, value in facts if label not in skip][:22]
    label_font, value_font = _font(22), _font(24, True)
    column_width = inner // 2
    row_height = 40
    for index, (label, value) in enumerate(grid):
        column, line = index % 2, index // 2
        x, top = margin + column * column_width, y + line * row_height
        if line % 2 == 0:
            draw.rectangle([x, top, x + column_width - 8, top + row_height - 2], fill="#f9fafb")
        draw.text((x + 12, top + 7), label, font=label_font, fill="#667085")
        shown = value
        while draw.textlength(shown, font=value_font) > column_width - 270 and len(shown) > 4:
            shown = shown[:-2]
        draw.text((x + 250, top + 5), shown + ("…" if shown != value else ""), font=value_font, fill="#101828")
    y += ((len(grid) + 1) // 2) * row_height + 14

    footer_top = height - 90
    body_font = _font(25)
    if data.get("Adres"):
        for line in _wrap_text(draw, "Adres: " + data["Adres"], body_font, inner)[:2]:
            draw.text((margin, y), line, font=body_font, fill="#344054")
            y += 34
        y += 6
    description = data.get("Açıklama")
    if description and y < footer_top - 80:
        draw.text((margin, y), "Açıklama", font=_font(28, True), fill="#101828")
        y += 42
        lines = _wrap_text(draw, description, body_font, inner)
        capacity = max(0, (footer_top - 12 - y) // 34)
        shown_lines = lines[:capacity]
        if len(lines) > capacity and shown_lines:
            shown_lines[-1] = shown_lines[-1].rstrip()[:-3] + "…"
        for line in shown_lines:
            draw.text((margin, y), line, font=body_font, fill="#344054")
            y += 34

    draw.line([margin, footer_top, width - margin, footer_top], fill="#d0d5dd", width=2)
    footer = f"Sorumlu: {data.get('Sorumlu', '—')}   ·   Hazırlanma tarihi: {date.today().strftime('%d.%m.%Y')}"
    footer_font = _font(24)
    draw.text(((width - draw.textlength(footer, font=footer_font)) / 2, footer_top + 26), footer, font=footer_font, fill="#667085")
    return page


def _make_export(kind: str, facts: list[tuple[str, str]], photo_urls: list[str], office_name: str = DEFAULT_OFFICE_NAME) -> bytes:
    images = _fetch_images([{"url": url} for url in photo_urls])
    if kind == "xlsx":
        return build_excel(facts, images, office_name)
    page = render_listing_image(facts, images, office_name)
    buffer = io.BytesIO()
    if kind == "png":
        page.save(buffer, "PNG")
    else:
        page.save(buffer, "PDF", resolution=150.0)
    return buffer.getvalue()


@st.cache_data(ttl=900, show_spinner="Dosya hazırlanıyor…", max_entries=30)
def _make_export_cached(kind: str, facts_key: tuple, photo_names: tuple, office_name: str, _photo_urls: tuple) -> bytes:
    return _make_export(kind, list(facts_key), list(_photo_urls), office_name)


def download_tab_content(row: pd.Series, photos: list[dict]) -> None:
    property_id = str(row.id)
    facts = property_fact_rows(row)
    urls = tuple(p["url"] for p in photos[:4])
    names = tuple(p["name"] for p in photos[:4])
    office_name = st.session_state.get("office_name", DEFAULT_OFFICE_NAME)
    specs = [
        ("xlsx", "📊 Excel indir", "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("pdf", "📄 PDF indir", "pdf", "application/pdf"),
        ("png", "🖼️ Resim indir", "png", "image/png"),
    ]
    for column, (kind, label, extension, mime) in zip(st.columns(3), specs):
        file_name = f"ilan_{property_id[:8]}.{extension}"
        if DOWNLOAD_CALLABLE:
            column.download_button(label, data=partial(_make_export, kind, facts, list(urls), office_name), file_name=file_name, mime=mime, key=f"dl_{kind}_{property_id}", use_container_width=True, on_click="ignore")
        else:
            try:
                data = _make_export_cached(kind, tuple(facts), names, office_name, urls)
            except Exception as exc:
                column.error(f"{label}: hazırlanamadı ({exc})")
                continue
            column.download_button(label, data, file_name=file_name, mime=mime, key=f"dl_{kind}_{property_id}", use_container_width=True)
    st.caption("Dosya, butona dokunduğunuzda hazırlanır ve iner; fotoğraf sayısına göre birkaç saniye sürebilir.")
    st.divider()
    whatsapp_block(row, facts, photos, office_name)


def add_request(client: Client, user_id: str, staff: pd.DataFrame, location_data: tuple[list[str], dict[str, list[str]]]) -> None:
    provinces, districts_by_province = location_data
    with st.container(border=True):
        st.subheader("Yeni müşteri talebi ekle")
        c1, c2, c3 = st.columns(3)
        customer = c1.text_input("Müşteri adı *")
        phone = c2.text_input("Telefon")
        deal = c3.selectbox("İşlem", ["Kiralık", "Satılık"])
        c4, c5, c6 = st.columns(3)
        city = c4.selectbox("İstenen il", ["Farketmez"] + provinces, key="request_city") if provinces else c4.text_input("İstenen il")
        district_options = districts_by_province.get(city, []) if city != "Farketmez" else []
        if city == "Farketmez":
            district = c5.selectbox("İstenen ilçe", ["Farketmez"], key="request_district_any")
        elif district_options:
            district = c5.selectbox("İstenen ilçe", district_options, key=f"request_district_{city}")
        else:
            district = c5.text_input("İstenen ilçe", key=f"request_district_text_{city}")
        rooms = c6.selectbox("Oda tercihi", ["Farketmez", "1+0", "1+1", "2+1", "3+1", "4+1", "Diğer"])
        c7, c8, c9 = st.columns(3)
        min_price = c7.number_input("En düşük bütçe (₺)", min_value=0, step=1000)
        max_price = c8.number_input("En yüksek bütçe (₺)", min_value=0, step=1000)
        deadline = c9.date_input("Hedef tarih", value=None)
        owner = st.selectbox("Sorumlu çalışan", staff["id"].tolist(), format_func=lambda x: staff.loc[staff.id == x, "name"].iloc[0], key="request_owner")
        notes = st.text_area("Müşteri notları")
        saved = st.button("Talebi kaydet", type="primary", use_container_width=True, key="save_new_request")
    if saved:
        if not customer:
            st.error("Müşteri adı alanını doldurun.")
            return
        if max_price and min_price > max_price:
            st.error("En düşük bütçe en yüksek bütçeden fazla olamaz.")
            return
        client.table("requests").insert({"customer_name": customer, "phone": phone, "deal_type": deal, "city": city, "district": district, "rooms": rooms, "min_price": int(min_price), "max_price": int(max_price), "target_date": deadline.isoformat() if isinstance(deadline, date) else None, "status": "Yeni", "notes": notes, "owner_id": owner, "created_by": user_id}).execute()
        st.success("Talep kaydedildi.")
        st.rerun()


BACKUP_FORMAT = "ofis-portfoy-backup-v1"
BACKUP_ALLOWED_COLUMNS = {
    "title", "city", "district", "deal_type", "price", "rooms", "area_m2", "net_area_m2",
    "property_category", "building_age", "floor", "total_floors", "heating", "bathroom_count",
    "kitchen", "balcony", "elevator", "parking", "furnished", "usage_status", "within_complex",
    "complex_name", "monthly_dues", "title_deed_status", "loan_eligible", "swap_allowed",
    "address", "description", "maps_url", "latitude", "longitude", "status", "owner_id", "created_by",
}


def build_backup_zip(admin_client: Client, include_photos: bool, progress=None) -> bytes:
    all_properties = rows(admin_client, "properties")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("ilanlar.json", json.dumps({"format": BACKUP_FORMAT, "photos_included": include_photos, "properties": all_properties}, ensure_ascii=False, default=str, indent=2), compress_type=zipfile.ZIP_DEFLATED)
        if include_photos:
            bucket = admin_client.storage.from_(PHOTO_BUCKET)
            for index, item in enumerate(all_properties, start=1):
                property_id = str(item["id"])
                for obj in bucket.list(property_id, {"limit": 1000}) or []:
                    if not obj.get("id"):
                        continue
                    try:
                        archive.writestr(f"fotograflar/{property_id}/{obj['name']}", bucket.download(f"{property_id}/{obj['name']}"), compress_type=zipfile.ZIP_STORED)
                    except Exception:
                        continue
                if progress:
                    progress(index, len(all_properties))
    return buffer.getvalue()


def restore_backup(admin_client: Client, uploaded, progress=None) -> tuple[int, int, list[str]]:
    content = uploaded.getvalue()
    archive = None
    photos_by_old: dict[str, list[str]] = {}
    if uploaded.name.lower().endswith(".zip"):
        archive = zipfile.ZipFile(io.BytesIO(content))
        data = json.loads(archive.read("ilanlar.json").decode("utf-8"))
        for entry in archive.namelist():
            parts = entry.split("/")
            if len(parts) == 3 and parts[0] == "fotograflar" and parts[2]:
                photos_by_old.setdefault(parts[1], []).append(entry)
    else:
        data = json.loads(content.decode("utf-8"))
    backup_rows = data.get("properties", []) if isinstance(data, dict) else []
    bucket = admin_client.storage.from_(PHOTO_BUCKET)
    mime_types = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}
    restored, photo_count, errors = 0, 0, []
    for number, row in enumerate(backup_rows, start=1):
        try:
            result = admin_client.table("properties").insert({k: v for k, v in row.items() if k in BACKUP_ALLOWED_COLUMNS}).execute()
            new_id = str(result.data[0]["id"])
            restored += 1
        except Exception as exc:
            errors.append(f"{number}. ilan ({row.get('title') or 'başlıksız'}): {exc}")
            continue
        for entry in photos_by_old.get(str(row.get("id")), []):
            extension = Path(entry).suffix.lower()
            try:
                bucket.upload(f"{new_id}/{uuid4().hex}{extension}", archive.read(entry), file_options={"content-type": mime_types.get(extension, "application/octet-stream"), "upsert": "false"})
                photo_count += 1
            except Exception as exc:
                errors.append(f"{number}. ilanın fotoğrafı yüklenemedi: {exc}")
        if progress:
            progress(number, len(backup_rows))
    return restored, photo_count, errors


def property_backup_tools(admin_client: Client) -> None:
    notice = st.session_state.pop("restore_notice", None)
    if notice:
        st.success(notice["message"])
        for error in notice["errors"]:
            st.error(error)

    st.markdown("#### İlanları yedekle")
    include_photos = st.checkbox("Fotoğrafları da yedeğe ekle", value=True, key="backup_include_photos")
    st.caption("Yedek tek bir ZIP dosyasıdır: ilan bilgileri ve (seçiliyse) tüm fotoğraflar. Çok fotoğraf varsa hazırlanması birkaç dakika sürebilir.")
    if st.button("Yedeği hazırla", key="prepare_backup"):
        bar = st.progress(0.0, text="Yedek hazırlanıyor…")
        try:
            data = build_backup_zip(admin_client, include_photos, lambda done, total: bar.progress(done / max(total, 1), text=f"{done}/{total} ilan"))
            st.session_state["backup_zip"] = (data, f"ilan_yedegi_{date.today().isoformat()}.zip")
        except Exception as exc:
            st.error("Yedek hazırlanamadı.")
            st.caption(str(exc))
        bar.empty()
    prepared = st.session_state.get("backup_zip")
    if prepared:
        st.download_button(f"⬇️ Yedeği indir ({len(prepared[0]) / 1_048_576:.1f} MB)", prepared[0], prepared[1], "application/zip", key="download_backup_zip", type="primary")

    st.divider()
    st.markdown("#### Yedek dosyasından ilanları geri yükle")
    st.warning("Geri yükleme, dosyadaki ilanları ve fotoğraflarını yeni kayıt olarak ekler; mevcut ilanları silmez veya güncellemez. Aynı yedeği iki kez yüklerseniz ilanlar çift olur.")
    backup_file = st.file_uploader("Yedek dosyası (ZIP veya eski JSON)", type=["zip", "json"], key="property_backup_upload")
    if backup_file and st.button("Yedeği geri yükle", type="primary", key="restore_property_backup"):
        bar = st.progress(0.0, text="Geri yükleniyor…")
        try:
            restored, photo_count, errors = restore_backup(admin_client, backup_file, lambda done, total: bar.progress(done / max(total, 1), text=f"{done}/{total} ilan"))
        except Exception as exc:
            st.error(f"Yedek dosyası okunamadı: {exc}")
            return
        st.session_state.restore_notice = {"message": f"{restored} ilan ve {photo_count} fotoğraf geri yüklendi.", "errors": errors}
        st.rerun()


def property_delete_tools(admin_client: Client, staff: pd.DataFrame) -> None:
    notice = st.session_state.pop("delete_notice", None)
    if notice:
        st.success(notice["message"])
        for error in notice["errors"]:
            st.error(error)
    try:
        all_properties = rows(admin_client, "properties")
    except Exception as exc:
        st.error("İlanlar okunamadı.")
        st.caption(str(exc))
        return
    if not all_properties:
        st.info("Silinecek ilan yok.")
        return
    generation = st.session_state.get("delete_generation", 0)
    names = dict(zip(staff.id.astype(str), staff.name))
    labels = {
        str(p["id"]): f"{p.get('title') or 'Başlıksız ilan'} — {p.get('city', '')}/{p.get('district', '')} — {p.get('status', '')} — {names.get(str(p.get('owner_id')), 'Atanmamış')}"
        for p in all_properties
    }

    def finish(deleted: int, errors: list[str]) -> None:
        st.session_state["delete_generation"] = generation + 1
        st.session_state["delete_notice"] = {"message": f"{deleted} ilan silindi.", "errors": errors}
        st.session_state.pop("selected_property_id", None)
        st.rerun()

    st.markdown("#### Seçili ilanları sil")
    selected = st.multiselect("Silinecek ilanlar", list(labels), format_func=labels.get, key=f"admin_delete_select_{generation}")
    confirm = st.checkbox("Seçili ilanların ve fotoğraflarının kalıcı olarak silineceğini anlıyorum.", key=f"admin_delete_confirm_{generation}")
    if st.button(f"Seçili {len(selected)} ilanı sil", type="primary", disabled=not (selected and confirm), key=f"admin_delete_button_{generation}"):
        finish(*delete_properties(admin_client, selected))

    st.divider()
    st.markdown("#### Tüm ilanları sil")
    st.error(f"Bu işlem {len(all_properties)} ilanın tamamını ve fotoğraflarını kalıcı olarak siler. Önce 'İlan yedeği' sekmesinden yedek almanız önerilir.")
    phrase = st.text_input("Onaylamak için TÜMÜNÜ SİL yazın", key=f"admin_delete_all_phrase_{generation}")
    if st.button("Tüm ilanları sil", type="primary", disabled=_norm(phrase) != _norm("TÜMÜNÜ SİL"), key=f"admin_delete_all_button_{generation}"):
        finish(*delete_properties(admin_client, list(labels)))


def switch_to_tab(name: str, click_tab: bool = True) -> None:
    nonce = st.session_state.get("goto_nonce", 0)
    script = """<script>
// __NONCE__
const wanted = __NAME__;
const clickTab = __CLICK__;
const W = window.parent;
let tries = 0;
function press(el) {
  ["pointerdown", "mousedown", "pointerup", "mouseup", "click"].forEach(type => {
    const Ctor = type.startsWith("pointer") && W.PointerEvent ? W.PointerEvent : W.MouseEvent;
    el.dispatchEvent(new Ctor(type, {bubbles: true, cancelable: true, view: W, button: 0}));
  });
}
function collapseSidebar(doc) {
  try {
    if (W.innerWidth < 800) {
      const collapse = doc.querySelector('[data-testid="stSidebarCollapseButton"] button');
      if (collapse) collapse.click();
    }
    W.scrollTo(0, 0);
  } catch (e) {}
}
const timer = setInterval(() => {
  tries += 1;
  let doc;
  try { doc = W.document; } catch (e) { clearInterval(timer); return; }
  if (!clickTab) { collapseSidebar(doc); clearInterval(timer); return; }
  const tab = Array.from(doc.querySelectorAll('[role="tab"]')).find(b => b.textContent.trim() === wanted);
  if (tab) {
    press(tab);
    clearInterval(timer);
    collapseSidebar(doc);
  } else if (tries > 30) { clearInterval(timer); }
}, 100);
</script>""".replace("__NONCE__", str(nonce)).replace("__NAME__", json.dumps(name)).replace("__CLICK__", "true" if click_tab else "false")
    embed_html(script, 1)


def go_to_section(name: str) -> None:
    st.session_state["pending_tab"] = name
    st.session_state["goto_nonce"] = st.session_state.get("goto_nonce", 0) + 1
    if name == "Daireler":
        st.session_state.pop("selected_property_id", None)


def admin_panel(client: Client, admin_client: Client, staff: pd.DataFrame, user_id: str) -> None:
    st.subheader("Çalışan hesapları")
    st.caption("Yeni kullanıcı oluşturabilir, çalışan bilgilerini düzenleyebilir veya hesabı geçici olarak devre dışı bırakabilirsiniz.")
    try:
        listed_users = admin_client.auth.admin.list_users()
        auth_users = listed_users if isinstance(listed_users, list) else listed_users.users
        user_emails = {str(u.id): u.email for u in auth_users}
    except Exception as exc:
        st.error("Auth kullanıcı listesi okunamadı. Admin key ayarını kontrol edin.")
        st.code(str(exc))
        return

    roster = staff.copy()
    roster["E-posta"] = roster["id"].astype(str).map(user_emails).fillna("—")
    roster["Durum"] = roster.get("is_active", True).apply(lambda active: "Aktif" if active else "Devre dışı") if isinstance(roster.get("is_active", True), pd.Series) else "Aktif"
    roster["Rol"] = roster["is_admin"].map(lambda admin: "Yönetici" if admin else "Çalışan")
    st.dataframe(roster[["name", "E-posta", "Rol", "Durum"]].rename(columns={"name": "Ad Soyad"}), hide_index=True, use_container_width=True)

    create_tab, edit_tab, backup_tab, delete_tab, titles_tab, news_tab = st.tabs(["Kullanıcı ekle", "Düzenle / devre dışı bırak", "İlan yedeği", "İlan sil", "Başlıklar", "Duyurular"])
    with backup_tab:
        property_backup_tools(admin_client)
    with delete_tab:
        property_delete_tools(admin_client, staff)
    with titles_tab:
        settings_admin_tab(client, admin_client)
    with news_tab:
        announcements_admin_tab(admin_client)
    with create_tab:
        with st.form("create_employee", clear_on_submit=True):
            name = st.text_input("Ad soyad")
            email = st.text_input("E-posta")
            password = st.text_input("İlk şifre", type="password", help="Güçlü ve geçici bir şifre belirleyip çalışana güvenli şekilde iletin.")
            role = st.selectbox("Yetki", ["Çalışan", "Yönetici"])
            create = st.form_submit_button("Kullanıcı oluştur", type="primary")
        if create:
            if not name.strip() or not email.strip() or len(password) < 8:
                st.error("Ad, e-posta girin ve en az 8 karakterli bir ilk şifre belirleyin.")
            else:
                created = None
                try:
                    result = admin_client.auth.admin.create_user({"email": email.strip(), "password": password, "email_confirm": True, "user_metadata": {"name": name.strip(), "display_name": name.strip()}})
                    created = result.user
                    admin_client.table("profiles").insert({"id": created.id, "name": name.strip(), "is_admin": role == "Yönetici", "is_active": True}).execute()
                    st.success(f"{email.strip()} hesabı oluşturuldu.")
                    st.rerun()
                except Exception as exc:
                    if created:
                        try:
                            admin_client.auth.admin.delete_user(str(created.id))
                        except Exception:
                            pass
                    st.error("Kullanıcı oluşturulamadı. E-posta daha önce kullanılmış olabilir veya Admin key yetkisini kontrol edin.")
                    st.code(str(exc))

    with edit_tab:
        if staff.empty:
            st.info("Düzenlenecek kullanıcı yok.")
            return
        ids = staff.id.astype(str).tolist()
        selected_id = st.selectbox("Çalışan", ids, format_func=lambda ident: f"{staff.loc[staff.id.astype(str) == ident, 'name'].iloc[0]} — {user_emails.get(ident, 'e-posta yok')}")
        current = staff[staff.id.astype(str) == selected_id].iloc[0]
        with st.form("edit_employee"):
            updated_name = st.text_input("Ad soyad", value=str(current.name))
            updated_email = st.text_input("E-posta", value=user_emails.get(selected_id, ""))
            updated_admin = st.checkbox("Yönetici yetkisi", value=bool(current.is_admin))
            active_default = bool(current.get("is_active", True))
            updated_active = st.checkbox("Hesap aktif", value=active_default)
            save = st.form_submit_button("Değişiklikleri kaydet", type="primary")
        if save:
            admins = staff[staff.is_admin.astype(bool)]
            active_admins = admins[admins.get("is_active", pd.Series(True, index=admins.index)).astype(bool)]
            if (not updated_admin or not updated_active) and selected_id in active_admins.id.astype(str).tolist() and len(active_admins) <= 1:
                st.error("Son aktif yöneticinin yetkisi kaldırılamaz veya hesabı kapatılamaz.")
            elif not updated_name.strip() or not updated_email.strip():
                st.error("Ad ve e-posta boş bırakılamaz.")
            else:
                try:
                    if selected_id == user_id and not updated_active:
                        st.error("Kendi hesabınızı buradan devre dışı bırakamazsınız.")
                        return
                    admin_client.auth.admin.update_user_by_id(selected_id, {"email": updated_email.strip(), "user_metadata": {"name": updated_name.strip(), "display_name": updated_name.strip()}})
                    admin_client.table("profiles").update({"name": updated_name.strip(), "is_admin": updated_admin, "is_active": updated_active}).eq("id", selected_id).execute()
                    st.success("Çalışan bilgileri güncellendi.")
                    st.rerun()
                except Exception as exc:
                    st.error("Değişiklik kaydedilemedi. Auth veya tablo izinlerini kontrol edin.")
                    st.code(str(exc))

        st.warning("Devre dışı bırakılan çalışan yeni oturum açamaz; mevcut daire ve talepleri korunur. Hesabı daha sonra yeniden etkinleştirebilirsiniz.")


OFFERS_SQL = """create table if not exists public.offers (
  id uuid primary key default gen_random_uuid(),
  property_id uuid not null references public.properties(id) on delete cascade,
  customer_name text not null,
  phone text,
  amount numeric not null default 0,
  status text not null default 'Beklemede',
  notes text,
  offer_date date default current_date,
  created_by uuid default auth.uid(),
  created_at timestamptz not null default now(),
  updated_at timestamptz not null default now()
);
create index if not exists offers_property_id_idx on public.offers(property_id);
alter table public.offers enable row level security;
grant select, insert, update, delete on public.offers to authenticated;
create policy "offers_select" on public.offers for select to authenticated using (true);
create policy "offers_insert" on public.offers for insert to authenticated with check (auth.uid() is not null);
create policy "offers_update" on public.offers for update to authenticated using (
  created_by = auth.uid()
  or exists (select 1 from public.profiles p where p.id = auth.uid() and p.is_admin)
  or exists (select 1 from public.properties pr where pr.id = property_id and pr.owner_id = auth.uid()));
create policy "offers_delete" on public.offers for delete to authenticated using (
  created_by = auth.uid()
  or exists (select 1 from public.profiles p where p.id = auth.uid() and p.is_admin)
  or exists (select 1 from public.properties pr where pr.id = property_id and pr.owner_id = auth.uid()));"""


def money_text(value) -> str:
    try:
        return f"{int(float(value)):,}".replace(",", ".")
    except (TypeError, ValueError):
        return "0"


# ---------- küçük resimler ----------
@st.cache_resource
def thumb_store() -> dict:
    return {}


def _make_thumb(content: bytes) -> bytes:
    from PIL import Image, ImageOps

    image = ImageOps.exif_transpose(Image.open(io.BytesIO(content))).convert("RGB")
    image = ImageOps.fit(image, (160, 120), method=Image.LANCZOS)
    buffer = io.BytesIO()
    image.save(buffer, "JPEG", quality=72)
    return buffer.getvalue()


def _fetch_thumb(client: Client, store: dict, property_id: str) -> None:
    try:
        bucket = client.storage.from_(PHOTO_BUCKET)
        objects = bucket.list(property_id, {"limit": 5, "sortBy": {"column": "name", "order": "asc"}}) or []
        first = next((o for o in objects if o.get("id")), None)
        data = _make_thumb(bucket.download(f"{property_id}/{first['name']}")) if first else None
        store[property_id] = (time.time() + 900, data)
    except Exception:
        store[property_id] = (time.time() + 60, None)


def ensure_thumbs(client: Client, property_ids: list[str], budget: float = 6.0, limit: int = 60) -> None:
    store = thumb_store()
    now = time.time()
    missing = [pid for pid in property_ids[:limit] if pid not in store or store[pid][0] < now]
    if not missing:
        return
    pool = ThreadPoolExecutor(max_workers=8)
    try:
        futures = [pool.submit(_fetch_thumb, client, store, pid) for pid in missing]
        wait_futures(futures, timeout=budget)
    finally:
        pool.shutdown(wait=False)


def open_employee(employee_id: str) -> None:
    st.session_state["summary_employee"] = str(employee_id)
    st.session_state["pending_tab"] = "Çalışan özeti"
    st.session_state["goto_nonce"] = st.session_state.get("goto_nonce", 0) + 1


def show_deal(deal: str) -> None:
    st.session_state["list_deal_filter"] = deal
    st.session_state["daire_view"] = "Tüm ilanlar"
    go_to_section("Daireler")


def info_tiles_html(row: pd.Series) -> str:
    def value(name: str) -> str:
        text = _clean(row.get(name))
        return "" if text in ("Belirtilmemiş", "Belirtilmedi", "0") and name != "building_age" else text

    def flag(name: str) -> bool:
        item = row.get(name)
        return bool(item) and not (isinstance(item, float) and pd.isna(item))

    dues = pd.to_numeric(row.get("monthly_dues"), errors="coerce")
    floor, age = value("floor"), value("building_age")
    rooms = value("rooms")
    tiles = [
        ("🚪", "Oda Sayısı", rooms if (not rooms or "Oda" in rooms) else f"{rooms} Oda"),
        ("📐", "Brüt Alan", f"{value('area_m2')} m²" if value("area_m2") else ""),
        ("📏", "Net Alan", f"{value('net_area_m2')} m²" if value("net_area_m2") else ""),
        ("🏗️", "Bina Yaşı", ("0 (Yeni)" if age == "0" else f"{age} Yıl") if age else ""),
        ("🏢", "Bulunduğu Kat", (f"{floor}. Kat" if floor.isdigit() else floor) if floor else ""),
        ("🔥", "Isıtma", value("heating")),
        ("🚿", "Banyo", value("bathroom_count")),
        ("🍳", "Mutfak", value("kitchen")),
        ("🚗", "Otopark", value("parking") if value("parking") != "Yok" else ""),
        ("🌇", "Balkon", "Var" if flag("balcony") else ""),
        ("🛗", "Asansör", "Var" if flag("elevator") else ""),
        ("🛋️", "Eşya Durumu", "Eşyalı" if flag("furnished") else ""),
        ("🏘️", "Site", (value("complex_name") or "Site içinde") if flag("within_complex") else ""),
        ("💳", "Aylık Aidat", f"{money_text(dues)} ₺" if pd.notna(dues) and dues > 0 else ""),
        ("📜", "Tapu Durumu", value("title_deed_status")),
        ("🏦", "Kredi", "Krediye Uygun" if flag("loan_eligible") else ""),
        ("🔄", "Takas", "Değerlendirilir" if flag("swap_allowed") else ""),
    ]
    cells = "".join(f"<div class='info-tile'><div class='it-icon'>{icon}</div><div class='it-label'>{escape(label)}</div><div class='it-value'>{escape(text)}</div></div>" for icon, label, text in tiles if text)
    return f"<div class='info-tiles'>{cells}</div>" if cells else ""


def open_property(property_id: str) -> None:
    st.session_state["selected_property_id"] = str(property_id)
    st.session_state["pending_tab"] = "Daireler"
    st.session_state["goto_nonce"] = st.session_state.get("goto_nonce", 0) + 1


def render_property_rows(client: Client, items: list[dict], scope: str, notes: dict[str, str] | None = None, height: int = 470, favorites: set[str] | None = None, row_action=None) -> None:
    ensure_thumbs(client, [str(item["id"]) for item in items])
    store = thumb_store()
    with st.container(height=height, border=False):
        if not items:
            st.info("Gösterilecek ilan yok.")
        for item in items:
            property_id = str(item["id"])
            area = _clean(item.get("area_m2")) or "—"
            line = " · ".join(part for part in [
                _clean(item.get("title")) or "Başlıksız ilan", _clean(item.get("deal_type")),
                f"{_clean(item.get('city'))} / {_clean(item.get('district'))}", _clean(item.get("rooms")),
                f"{area} m²", money_text(item.get("price")) + " ₺", _clean(item.get("status")), _clean(item.get("sorumlu")),
            ] if part)
            if favorites and property_id in favorites:
                line = "⭐ " + line
            thumb = (store.get(property_id) or (0, None))[1]
            with st.container(key=f"listing_row_{scope}_{property_id}"):
                image_col, button_col = st.columns([1, 9], vertical_alignment="center")
                with image_col:
                    st.image(thumb if thumb else str(DEFAULT_PROPERTY_IMAGE), width=64)
                with button_col:
                    st.button(line, key=f"open_{scope}_{property_id}", use_container_width=True, on_click=open_property, args=(property_id,))
                    if notes and property_id in notes:
                        st.caption(notes[property_id])
                    if row_action:
                        row_action(item, property_id)


# ---------- müşteri talebi eşleştirme ----------
def _room_key(value) -> str:
    text = _clean(value)
    found = re.search(r"\d\+\d", text)
    return found.group(0) if found else _norm(text)


def match_properties(request: dict, properties: pd.DataFrame, min_score: int = 50) -> list[tuple[dict, int, str]]:
    if properties.empty:
        return []
    deal = _clean(request.get("deal_type"))
    city_req, district_req, rooms_req = _clean(request.get("city")), _clean(request.get("district")), _clean(request.get("rooms"))
    low, high = int(request.get("min_price") or 0), int(request.get("max_price") or 0)
    wants = lambda text: _norm(text) not in ("", "farketmez", "belirtilmedi")
    results = []
    for item in properties.to_dict("records"):
        if _clean(item.get("status")) != "Aktif":
            continue
        if deal and _clean(item.get("deal_type")) != deal:
            continue
        earned, total, notes = 0, 0, []
        if deal:
            total += 30
            earned += 30
            notes.append(f"✅ {deal}")
        if wants(city_req):
            total += 10
            if _norm(item.get("city")) == _norm(city_req):
                earned += 10
                notes.append(f"✅ {city_req}")
            else:
                notes.append("❌ İl")
        if wants(district_req):
            total += 15
            if _norm(item.get("district")) == _norm(district_req):
                earned += 15
                notes.append(f"✅ {district_req}")
            else:
                notes.append("❌ İlçe")
        if wants(rooms_req):
            total += 15
            if _room_key(item.get("rooms")) == _room_key(rooms_req):
                earned += 15
                notes.append(f"✅ {rooms_req}")
            else:
                notes.append("❌ Oda")
        if low or high:
            total += 30
            price = int(float(item.get("price") or 0))
            if (not low or price >= low) and (not high or price <= high):
                earned += 30
                notes.append("✅ Bütçeye uygun")
            else:
                limit = high if (high and price > high) else low
                gap = abs(price - limit) / max(limit, 1)
                if gap <= 0.10:
                    earned += 12
                    notes.append(f"⚠️ Bütçenin %{gap * 100:.0f} dışında")
                else:
                    notes.append("❌ Bütçe dışı")
        score = round(100 * earned / total) if total else 100
        if score >= min_score:
            results.append((item, score, " · ".join(notes)))
    results.sort(key=lambda entry: (-entry[1], int(float(entry[0].get("price") or 0))))
    return results


def select_request(request_id: str) -> None:
    st.session_state["selected_request_id"] = str(request_id)


def requests_tab(client: Client, requests_df: pd.DataFrame, properties: pd.DataFrame, staff: pd.DataFrame, location_data: tuple[list[str], dict[str, list[str]]], user_id: str, is_admin: bool, favorites: set[str]) -> None:
    if requests_df.empty:
        st.info("Henüz müşteri talebi yok.")
        return
    items = requests_df.to_dict("records")
    if st.toggle("Tablo görünümü", key="requests_table_view"):
        table = pd.DataFrame([{
            "Müşteri": _clean(i.get("customer_name")), "Telefon": _clean(i.get("phone")), "İşlem": _clean(i.get("deal_type")),
            "İl": _clean(i.get("city")), "İlçe": _clean(i.get("district")), "Oda": _clean(i.get("rooms")),
            "En düşük bütçe (₺)": money_text(i.get("min_price")), "En yüksek bütçe (₺)": money_text(i.get("max_price")),
            "Hedef tarih": _clean(i.get("target_date")), "Durum": _clean(i.get("status")), "Sorumlu": _clean(i.get("sorumlu")), "Notlar": _clean(i.get("notes")),
        } for i in items])
        st.dataframe(table, hide_index=True, use_container_width=True, height=460)
        return
    provinces, districts_by_province = location_data
    ids = [str(item["id"]) for item in items]
    selected_id = st.session_state.get("selected_request_id")
    if selected_id not in ids:
        selected_id = ids[0]
    request = items[ids.index(selected_id)]
    can_manage = is_admin or str(request.get("owner_id")) == user_id or str(request.get("created_by")) == user_id
    left, right = st.columns([1, 2], gap="medium")
    with left:
        st.caption(f"{len(items)} talep · incelemek için dokunun")
        with st.container(height=520, border=False):
            for item in items:
                request_id = str(item["id"])
                budget = f"₺{money_text(item.get('min_price'))}–{money_text(item.get('max_price'))}" if (item.get("min_price") or item.get("max_price")) else ""
                line = " · ".join(part for part in [
                    f"👤 {_clean(item.get('customer_name'))}", _clean(item.get("deal_type")),
                    f"{_clean(item.get('city'))}/{_clean(item.get('district'))}", _clean(item.get("rooms")), budget, _clean(item.get("status")),
                ] if part)
                st.button(line, key=f"request_{request_id}", use_container_width=True, type="primary" if request_id == selected_id else "secondary", on_click=select_request, args=(request_id,))
    with right:
        with st.container(border=True):
            st.markdown(f"**👤 {_clean(request.get('customer_name'))}** · {_clean(request.get('phone')) or 'Telefon yok'} · Sorumlu: {_clean(request.get('sorumlu'))}")
            budget_text = f"₺{money_text(request.get('min_price'))} – ₺{money_text(request.get('max_price'))}" if (request.get("min_price") or request.get("max_price")) else "Belirtilmemiş"
            st.caption(f"{_clean(request.get('deal_type'))} · {_clean(request.get('city'))} / {_clean(request.get('district'))} · Oda: {_clean(request.get('rooms'))} · Bütçe: {budget_text} · Hedef tarih: {_clean(request.get('target_date')) or '—'}")
            if _clean(request.get("notes")):
                st.write(_clean(request.get("notes")))
            choices = ["Yeni", "İletişimde", "Gösterim", "Tamamlandı", "İptal"]
            current_status = _clean(request.get("status"))
            s1, s2 = st.columns([2, 1], vertical_alignment="bottom")
            new_status = s1.selectbox("Talep durumu", choices, index=choices.index(current_status) if current_status in choices else 0, key=f"request_status_{selected_id}")
            if s2.button("Durumu kaydet", key=f"request_status_save_{selected_id}", use_container_width=True):
                client.table("requests").update({"status": new_status}).eq("id", selected_id).execute()
                st.rerun()
        if can_manage:
            edit_col, delete_col = st.columns(2)
            with edit_col.expander("✏️ Talebi düzenle"):
                prefix = f"reqedit_{selected_id}"
                e1, e2, e3 = st.columns(3)
                customer = e1.text_input("Müşteri adı *", value=_clean(request.get("customer_name")), key=f"{prefix}_customer")
                phone = e2.text_input("Telefon", value=_clean(request.get("phone")), key=f"{prefix}_phone")
                deal_options = ["Kiralık", "Satılık"]
                deal = e3.selectbox("İşlem", deal_options, index=deal_options.index(_clean(request.get("deal_type"))) if _clean(request.get("deal_type")) in deal_options else 0, key=f"{prefix}_deal")
                e4, e5, e6 = st.columns(3)
                city_options = ["Farketmez"] + provinces
                current_city = _clean(request.get("city")) or "Farketmez"
                if provinces:
                    city = e4.selectbox("İstenen il", city_options, index=city_options.index(current_city) if current_city in city_options else 0, key=f"{prefix}_city")
                    district_options = ["Farketmez"] + (districts_by_province.get(city, []) if city != "Farketmez" else [])
                    current_district = _clean(request.get("district")) or "Farketmez"
                    district = e5.selectbox("İstenen ilçe", district_options, index=district_options.index(current_district) if current_district in district_options else 0, key=f"{prefix}_district_{city}")
                else:
                    city = e4.text_input("İstenen il", value=current_city, key=f"{prefix}_city")
                    district = e5.text_input("İstenen ilçe", value=_clean(request.get("district")), key=f"{prefix}_district")
                room_options = ["Farketmez", "1+0", "1+1", "2+1", "3+1", "4+1", "Diğer"]
                rooms = e6.selectbox("Oda tercihi", room_options, index=room_options.index(_clean(request.get("rooms"))) if _clean(request.get("rooms")) in room_options else 0, key=f"{prefix}_rooms")
                e7, e8, e9 = st.columns(3)
                min_price = e7.number_input("En düşük bütçe (₺)", min_value=0, step=1000, value=int(float(request.get("min_price") or 0)), key=f"{prefix}_min")
                max_price = e8.number_input("En yüksek bütçe (₺)", min_value=0, step=1000, value=int(float(request.get("max_price") or 0)), key=f"{prefix}_max")
                existing_date = pd.to_datetime(request.get("target_date"), errors="coerce")
                deadline = e9.date_input("Hedef tarih", value=existing_date.date() if pd.notna(existing_date) else None, key=f"{prefix}_date")
                owner_ids = staff["id"].astype(str).tolist()
                current_owner = str(request.get("owner_id"))
                owner = st.selectbox("Sorumlu çalışan", owner_ids, index=owner_ids.index(current_owner) if current_owner in owner_ids else 0, format_func=lambda x: staff.loc[staff.id.astype(str) == x, "name"].iloc[0], key=f"{prefix}_owner")
                notes = st.text_area("Müşteri notları", value=_clean(request.get("notes")), key=f"{prefix}_notes")
                if st.button("Değişiklikleri kaydet", type="primary", key=f"{prefix}_save", use_container_width=True):
                    if not customer.strip():
                        st.error("Müşteri adı boş olamaz.")
                    elif max_price and min_price > max_price:
                        st.error("En düşük bütçe en yüksek bütçeden fazla olamaz.")
                    else:
                        result = service_client(client).table("requests").update({
                            "customer_name": customer.strip(), "phone": phone.strip(), "deal_type": deal, "city": city, "district": district,
                            "rooms": rooms, "min_price": int(min_price), "max_price": int(max_price),
                            "target_date": deadline.isoformat() if isinstance(deadline, date) else None, "notes": notes.strip(), "owner_id": owner,
                        }).eq("id", selected_id).execute()
                        if result.data:
                            st.success("Talep güncellendi.")
                            st.rerun()
                        else:
                            st.error("Talep güncellenemedi.")
            with delete_col.expander("🗑️ Talebi sil"):
                st.warning("Talep kalıcı olarak silinir.")
                confirm = st.checkbox("Bu talebi silmek istediğimi onaylıyorum.", key=f"reqdel_confirm_{selected_id}")
                if st.button("Talebi sil", type="primary", disabled=not confirm, key=f"reqdel_{selected_id}", use_container_width=True):
                    result = service_client(client).table("requests").delete().eq("id", selected_id).execute()
                    if result.data:
                        st.session_state.pop("selected_request_id", None)
                        st.session_state["property_delete_notice"] = "Talep silindi."
                        st.rerun()
                    else:
                        st.error("Talep silinemedi.")
        m1, m2 = st.columns([2, 1], vertical_alignment="center")
        m1.markdown("**Bu talebe uygun ilanlar** (en uygun olan üstte)")
        min_score = m2.select_slider("En az uyum", options=[0, 30, 50, 70, 90], value=50, key="match_min_score", format_func=lambda v: f"%{v}")
        matches = match_properties(request, properties, int(min_score))
        if not matches:
            st.info("Bu talebe uygun aktif ilan bulunamadı. Uyum eşiğini düşürmeyi deneyin.")
        else:
            render_property_rows(client, [m[0] for m in matches], "match", notes={str(m[0]["id"]): f"%{m[1]} uyum · {m[2]}" for m in matches}, height=260, favorites=favorites)


# ---------- teklifler ----------
def load_offers(client: Client, property_id: str | None = None) -> list[dict]:
    query = client.table("offers").select("*").order("created_at", desc=True)
    if property_id:
        query = query.eq("property_id", property_id)
    return query.execute().data or []


def offers_setup_message(exc: Exception) -> None:
    st.warning("Teklifler için Supabase'de `offers` tablosu gerekli. Aşağıdaki SQL'i Supabase SQL Editor'da bir kez çalıştırın.")
    with st.expander("Çalıştırılacak SQL"):
        st.code(OFFERS_SQL, language="sql")
    st.caption(str(exc))


def offers_panel(client: Client, property_id: str, user_id: str, is_admin: bool, staff: pd.DataFrame, property_owner_id: str, scope: str) -> None:
    key = f"{scope}_{property_id}"
    generation = st.session_state.get(f"offer_generation_{key}", 0)

    def bump() -> None:
        st.session_state[f"offer_generation_{key}"] = generation + 1

    try:
        offers = load_offers(client, property_id)
    except Exception as exc:
        offers_setup_message(exc)
        return
    names = dict(zip(staff.id.astype(str), staff.name))
    if offers:
        top = max(float(o.get("amount") or 0) for o in offers)
        st.caption(f"{len(offers)} teklif · en yüksek teklif: {money_text(top)} ₺")
    with st.expander("➕ Yeni teklif ekle", expanded=not offers):
        a1, a2, a3 = st.columns(3)
        customer = a1.text_input("Müşteri adı *", key=f"offer_customer_{key}_{generation}")
        phone = a2.text_input("Telefon", key=f"offer_phone_{key}_{generation}")
        amount = a3.number_input("Teklif tutarı (₺) *", min_value=0, step=1000, key=f"offer_amount_{key}_{generation}")
        b1, b2 = st.columns([1, 2])
        status = b1.selectbox("Durum", OFFER_STATUSES, key=f"offer_status_{key}_{generation}")
        notes = b2.text_area("Açıklama", height=80, key=f"offer_notes_{key}_{generation}", placeholder="Ödeme şekli, şartlar, görüşme notu…")
        if st.button("Teklifi kaydet", type="primary", key=f"offer_save_{key}_{generation}"):
            if not customer.strip() or amount <= 0:
                st.error("Müşteri adı ve teklif tutarını girin.")
            else:
                try:
                    client.table("offers").insert({"property_id": property_id, "customer_name": customer.strip(), "phone": phone.strip(), "amount": int(amount), "status": status, "notes": notes.strip(), "created_by": user_id}).execute()
                    bump()
                    st.rerun()
                except Exception as exc:
                    st.error("Teklif kaydedilemedi.")
                    st.caption(str(exc))
    for offer in offers:
        offer_id = str(offer["id"])
        can_edit = is_admin or str(offer.get("created_by")) == user_id or str(property_owner_id) == user_id
        offer_date = pd.to_datetime(offer.get("offer_date") or offer.get("created_at"), errors="coerce")
        date_text = offer_date.strftime("%d.%m.%Y") if pd.notna(offer_date) else ""
        label = f"💰 {money_text(offer.get('amount'))} ₺ · {_clean(offer.get('customer_name'))} · {_clean(offer.get('status'))} · {date_text}"
        with st.expander(label):
            st.caption(f"Ekleyen: {names.get(str(offer.get('created_by')), '—')} · Telefon: {_clean(offer.get('phone')) or '—'}")
            if not can_edit:
                st.write(_clean(offer.get("notes")) or "Açıklama girilmemiş.")
                continue
            e1, e2, e3 = st.columns(3)
            new_customer = e1.text_input("Müşteri adı", value=_clean(offer.get("customer_name")), key=f"offer_edit_customer_{offer_id}_{generation}")
            new_phone = e2.text_input("Telefon", value=_clean(offer.get("phone")), key=f"offer_edit_phone_{offer_id}_{generation}")
            new_amount = e3.number_input("Tutar (₺)", min_value=0, step=1000, value=int(float(offer.get("amount") or 0)), key=f"offer_edit_amount_{offer_id}_{generation}")
            f1, f2 = st.columns([1, 2])
            current_status = _clean(offer.get("status"))
            options = OFFER_STATUSES if current_status in OFFER_STATUSES else OFFER_STATUSES + [current_status]
            new_status = f1.selectbox("Durum", options, index=options.index(current_status) if current_status in options else 0, key=f"offer_edit_status_{offer_id}_{generation}")
            new_notes = f2.text_area("Açıklama", value=_clean(offer.get("notes")), height=80, key=f"offer_edit_notes_{offer_id}_{generation}")
            g1, g2, g3 = st.columns([1, 1, 1], vertical_alignment="center")
            if g1.button("Güncelle", type="primary", key=f"offer_update_{offer_id}_{generation}", use_container_width=True):
                if not new_customer.strip() or new_amount <= 0:
                    st.error("Müşteri adı ve tutar boş olamaz.")
                else:
                    try:
                        result = client.table("offers").update({"customer_name": new_customer.strip(), "phone": new_phone.strip(), "amount": int(new_amount), "status": new_status, "notes": new_notes.strip(), "updated_at": pd.Timestamp.utcnow().isoformat()}).eq("id", offer_id).execute()
                        if result.data:
                            bump()
                            st.rerun()
                        else:
                            st.error("Teklif güncellenemedi (yetkiniz olmayabilir).")
                    except Exception as exc:
                        st.error("Teklif güncellenemedi.")
                        st.caption(str(exc))
            confirm = g2.checkbox("Silmeyi onaylıyorum", key=f"offer_confirm_{offer_id}_{generation}")
            if g3.button("Teklifi sil", disabled=not confirm, key=f"offer_delete_{offer_id}_{generation}", use_container_width=True):
                try:
                    result = client.table("offers").delete().eq("id", offer_id).execute()
                    if result.data:
                        bump()
                        st.rerun()
                    else:
                        st.error("Teklif silinemedi (yetkiniz olmayabilir).")
                except Exception as exc:
                    st.error("Teklif silinemedi.")
                    st.caption(str(exc))


def offers_tab(client: Client, properties: pd.DataFrame, staff: pd.DataFrame, user_id: str, is_admin: bool) -> None:
    if properties.empty:
        st.info("Teklif eklemek için önce bir ilan kaydedin.")
        return
    labels = {str(r["id"]): f"{_clean(r.get('title')) or 'Başlıksız'} — {_clean(r.get('city'))}/{_clean(r.get('district'))} — {_clean(r.get('deal_type'))} — {money_text(r.get('price'))} ₺" for r in properties.to_dict("records")}
    c1, c2 = st.columns([3, 1])
    choice = c1.selectbox("İlan seçin", ["Tüm ilanlar"] + list(labels), format_func=lambda x: x if x == "Tüm ilanlar" else labels[x], key="offers_property")
    if choice != "Tüm ilanlar":
        owner = str(properties.loc[properties.id.astype(str) == choice, "owner_id"].iloc[0])
        offers_panel(client, choice, user_id, is_admin, staff, owner, "main")
        return
    try:
        offers = load_offers(client)
    except Exception as exc:
        offers_setup_message(exc)
        return
    if not offers:
        st.info("Henüz teklif yok. Bir ilan seçip teklif ekleyebilirsiniz.")
        return
    names = dict(zip(staff.id.astype(str), staff.name))
    titles = {pid: label.split(" — ")[0] for pid, label in labels.items()}
    frame = pd.DataFrame([{
        "İlan": titles.get(str(o["property_id"]), "Silinmiş ilan"), "Müşteri": _clean(o.get("customer_name")), "Telefon": _clean(o.get("phone")),
        "Tutar (₺)": int(float(o.get("amount") or 0)), "Durum": _clean(o.get("status")),
        "Tarih": (pd.to_datetime(o.get("offer_date") or o.get("created_at"), errors="coerce")).strftime("%d.%m.%Y") if pd.notna(pd.to_datetime(o.get("offer_date") or o.get("created_at"), errors="coerce")) else "",
        "Ekleyen": names.get(str(o.get("created_by")), "—"), "Açıklama": _clean(o.get("notes")),
    } for o in offers])
    wanted = c2.multiselect("Durum", OFFER_STATUSES, key="offers_status_filter")
    if wanted:
        frame = frame[frame["Durum"].isin(wanted)]
    st.caption(f"{len(frame)} teklif · bir ilanı seçerek teklif ekleyebilir, düzenleyebilir veya silebilirsiniz.")
    st.dataframe(frame, hide_index=True, use_container_width=True, height=380)


# ---------- çalışan özeti ----------
# ---------- favoriler ----------
def load_favorites(client: Client, user_id: str) -> tuple[set[str], bool]:
    try:
        data = client.table("favorites").select("property_id").eq("user_id", user_id).execute().data or []
        return {str(item["property_id"]) for item in data}, True
    except Exception:
        return set(), False


def set_favorite(client: Client, user_id: str, property_id: str, add: bool) -> None:
    try:
        if add:
            client.table("favorites").upsert({"user_id": user_id, "property_id": property_id}).execute()
        else:
            client.table("favorites").delete().eq("user_id", user_id).eq("property_id", property_id).execute()
    except Exception as exc:
        st.session_state["feature_error"] = ("Favoriler", str(exc))


def favorites_view(client: Client, user_id: str, properties: pd.DataFrame, favorites: set[str]) -> None:
    items = properties[properties.id.astype(str).isin(favorites)].to_dict("records") if not properties.empty else []
    st.caption(f"{len(items)} favori ilan · ilana dokunarak açabilir, ★ ile favorilerden kaldırabilirsiniz.")

    def remove_action(item: dict, property_id: str) -> None:
        st.button("★ Favorilerden kaldır", key=f"unfav_{property_id}", on_click=set_favorite, args=(client, user_id, property_id, False))

    if not items:
        st.info("Henüz favori ilanınız yok. İlan detayında “Favorilere ekle” butonunu kullanın.")
        return
    render_property_rows(client, items, "fav", favorites=favorites, row_action=remove_action, height=440)


# ---------- fiyat geçmişi ----------
def load_price_history(client: Client, property_id: str | None = None, limit: int = 500) -> list[dict]:
    query = client.table("price_history").select("*").order("changed_at", desc=True).limit(limit)
    if property_id:
        query = query.eq("property_id", property_id)
    return query.execute().data or []


def _change_text(entry: dict) -> tuple[str, float]:
    old, new = float(entry.get("old_price") or 0), float(entry.get("new_price") or 0)
    pct = (new - old) / old * 100 if old else 0.0
    arrow = "📉" if new < old else "📈"
    stamp = pd.to_datetime(entry.get("changed_at"), utc=True, errors="coerce")
    when = stamp.tz_convert(TZ).strftime("%d.%m.%Y %H:%M") if pd.notna(stamp) else ""
    return f"{arrow} {money_text(old)} → {money_text(new)} ₺ ({pct:+.1f}%) · {when}", pct


def price_history_table(client: Client, property_id: str) -> None:
    try:
        history = load_price_history(client, property_id)
    except Exception as exc:
        feature_setup_message("Fiyat geçmişi", exc)
        return
    if not history:
        st.caption("Bu ilanda henüz fiyat değişikliği kaydı yok.")
        return
    names = st.session_state.get("staff_names", {})
    frame = pd.DataFrame([{
        "Tarih": (pd.to_datetime(h.get("changed_at"), utc=True).tz_convert(TZ).strftime("%d.%m.%Y %H:%M")),
        "Eski fiyat (₺)": money_text(h.get("old_price")), "Yeni fiyat (₺)": money_text(h.get("new_price")),
        "Değişim": _change_text(h)[0].split("(")[1].split(")")[0], "Değiştiren": names.get(str(h.get("changed_by")), "—"),
    } for h in history])
    st.dataframe(frame, hide_index=True, use_container_width=True)


def price_changes_view(client: Client, properties: pd.DataFrame) -> None:
    mode = st.radio("Fiyat değişimi", ["Tümü", "📉 Fiyatı düşenler", "📈 Fiyatı artanlar"], horizontal=True, key="price_change_filter", label_visibility="collapsed")
    try:
        history = load_price_history(client)
    except Exception as exc:
        feature_setup_message("Fiyat geçmişi", exc)
        return
    latest: dict[str, dict] = {}
    for entry in history:
        latest.setdefault(str(entry["property_id"]), entry)
    items, notes = [], {}
    for item in properties.to_dict("records") if not properties.empty else []:
        entry = latest.get(str(item["id"]))
        if not entry:
            continue
        text, pct = _change_text(entry)
        if mode.startswith("📉") and pct >= 0:
            continue
        if mode.startswith("📈") and pct <= 0:
            continue
        items.append((str(entry.get("changed_at")), item))
        notes[str(item["id"])] = text
    items.sort(key=lambda pair: pair[0], reverse=True)
    st.caption(f"{len(items)} ilanda fiyat değişikliği kaydı var.")
    if not items:
        st.info("Gösterilecek fiyat değişikliği yok. Fiyat değişiklikleri ilan düzenlendiğinde otomatik kaydedilir.")
        return
    render_property_rows(client, [pair[1] for pair in items], "price", notes=notes, height=420)


# ---------- mesajlar ----------
def load_messages(client: Client, user_id: str) -> list[dict]:
    return client.table("messages").select("*").or_(f"sender_id.eq.{user_id},recipient_id.eq.{user_id}").order("created_at", desc=True).limit(300).execute().data or []


def send_message(client: Client, sender_id: str, recipient_id: str, subject: str, body: str, property_id: str | None = None, parent_id: str | None = None) -> None:
    payload = {"sender_id": sender_id, "recipient_id": recipient_id, "subject": subject.strip(), "body": body.strip()}
    if property_id:
        payload["property_id"] = property_id
    if parent_id:
        payload["parent_id"] = parent_id
    client.table("messages").insert(payload).execute()


def compose_message_form(client: Client, user_id: str, staff: pd.DataFrame, scope: str, default_recipient: str | None = None, property_id: str | None = None, subject_hint: str = "", lock_recipient: bool = False) -> None:
    generation = st.session_state.get("msg_generation", 0)
    others = staff[staff.id.astype(str) != user_id]
    ids = others.id.astype(str).tolist()
    names = dict(zip(others.id.astype(str), others.name))
    if not ids:
        st.info("Mesaj gönderilebilecek başka çalışan yok.")
        return
    if lock_recipient and default_recipient in names:
        recipient = default_recipient
        st.caption(f"Alıcı: **{names[recipient]}**")
    else:
        recipient = st.selectbox("Alıcı", ids, index=ids.index(default_recipient) if default_recipient in ids else 0, format_func=names.get, key=f"msg_to_{scope}_{generation}")
    subject = st.text_input("Konu", value=subject_hint, key=f"msg_subject_{scope}_{generation}")
    body = st.text_area("Mesaj", height=100, key=f"msg_body_{scope}_{generation}")
    if st.button("Mesajı gönder", type="primary", key=f"msg_send_{scope}_{generation}"):
        if not body.strip():
            st.error("Mesaj metnini yazın.")
            return
        try:
            send_message(client, user_id, recipient, subject, body, property_id)
        except Exception as exc:
            feature_setup_message("Mesajlar", exc)
            return
        st.session_state["msg_generation"] = generation + 1
        st.session_state["message_notice"] = f"Mesaj {names.get(recipient, '')} kişisine gönderildi."
        st.rerun()


def messages_tab(client: Client, staff: pd.DataFrame, user_id: str, messages: list[dict], available: bool, properties: pd.DataFrame) -> None:
    if not available:
        feature_setup_message("Mesajlar")
        return
    names = dict(zip(staff.id.astype(str), staff.name))
    titles = {str(r["id"]): _clean(r.get("title")) for r in properties.to_dict("records")} if not properties.empty else {}
    inbox = [m for m in messages if str(m["recipient_id"]) == user_id and not m.get("deleted_by_recipient")]
    sent = [m for m in messages if str(m["sender_id"]) == user_id and not m.get("deleted_by_sender")]
    generation = st.session_state.get("msg_generation", 0)

    def bump() -> None:
        st.session_state["msg_generation"] = generation + 1

    def when(message: dict) -> str:
        stamp = pd.to_datetime(message.get("created_at"), utc=True, errors="coerce")
        return stamp.tz_convert(TZ).strftime("%d.%m.%Y %H:%M") if pd.notna(stamp) else ""

    def update(message_id: str, values: dict) -> None:
        client.table("messages").update(values).eq("id", message_id).execute()

    inbox_tab, sent_tab, new_tab = st.tabs(["Gelen kutusu", "Gönderilenler", "Yeni mesaj"])
    with inbox_tab:
        unread = [m for m in inbox if not m.get("is_read")]
        top1, top2 = st.columns([3, 1], vertical_alignment="center")
        top1.caption(f"{len(inbox)} mesaj · {len(unread)} okunmamış")
        if unread and top2.button("Tümünü okundu say", key=f"msg_read_all_{generation}", use_container_width=True):
            for message in unread:
                update(str(message["id"]), {"is_read": True})
            bump()
            st.rerun()
        if not inbox:
            st.info("Gelen kutunuz boş.")
        with st.container(height=430, border=False):
            for message in inbox:
                message_id = str(message["id"])
                sender = names.get(str(message["sender_id"]), "Bilinmiyor")
                mark = "🔵 " if not message.get("is_read") else ""
                with st.expander(f"{mark}{sender} · {message.get('subject') or '(konusuz)'} · {when(message)}"):
                    property_id = str(message.get("property_id") or "")
                    if property_id and property_id in titles:
                        st.button(f"🏠 İlanı aç: {titles[property_id]}", key=f"msg_open_{message_id}_{generation}", on_click=open_property, args=(property_id,))
                    st.write(message["body"])
                    reply = st.text_area("Yanıtınız", height=80, key=f"msg_reply_{message_id}_{generation}")
                    b1, b2, b3 = st.columns(3)
                    if b1.button("Yanıtla", type="primary", key=f"msg_reply_send_{message_id}_{generation}", use_container_width=True):
                        if not reply.strip():
                            st.error("Yanıt metnini yazın.")
                        else:
                            subject = message.get("subject") or ""
                            send_message(client, user_id, str(message["sender_id"]), subject if subject.startswith("Re:") else f"Re: {subject}".strip(), reply, property_id or None, message_id)
                            update(message_id, {"is_read": True})
                            bump()
                            st.session_state["message_notice"] = "Yanıt gönderildi."
                            st.rerun()
                    if not message.get("is_read") and b2.button("Okundu say", key=f"msg_read_{message_id}_{generation}", use_container_width=True):
                        update(message_id, {"is_read": True})
                        bump()
                        st.rerun()
                    if b3.button("Sil", key=f"msg_delete_in_{message_id}_{generation}", use_container_width=True):
                        update(message_id, {"deleted_by_recipient": True})
                        bump()
                        st.rerun()
    with sent_tab:
        st.caption(f"{len(sent)} gönderilmiş mesaj")
        if not sent:
            st.info("Henüz mesaj göndermediniz.")
        with st.container(height=430, border=False):
            for message in sent:
                message_id = str(message["id"])
                with st.expander(f"→ {names.get(str(message['recipient_id']), 'Bilinmiyor')} · {message.get('subject') or '(konusuz)'} · {when(message)}"):
                    st.write(message["body"])
                    st.caption("Okundu" if message.get("is_read") else "Henüz okunmadı")
                    if st.button("Sil", key=f"msg_delete_out_{message_id}_{generation}"):
                        update(message_id, {"deleted_by_sender": True})
                        bump()
                        st.rerun()
    with new_tab:
        compose_message_form(client, user_id, staff, "main")


# ---------- ajanda ----------
def load_reminders(client: Client, user_id: str) -> list[dict]:
    return client.table("reminders").select("*").eq("user_id", user_id).order("remind_at").execute().data or []


def reminder_time(reminder: dict) -> datetime:
    return pd.to_datetime(reminder["remind_at"], utc=True).tz_convert(TZ).to_pydatetime()


def reminder_alerts(reminders: list[dict]) -> tuple[list[dict], list[dict]]:
    now = now_tr()
    due = [r for r in reminders if not r.get("done") and reminder_time(r) <= now]
    today = [r for r in reminders if not r.get("done") and reminder_time(r) > now and reminder_time(r).date() == now.date()]
    return due, today


def agenda_tab(client: Client, user_id: str, properties: pd.DataFrame, reminders: list[dict], available: bool) -> None:
    if not available:
        feature_setup_message("Ajanda")
        return
    due, today = reminder_alerts(reminders)
    if due:
        st.error("⏰ **Zamanı gelen / geçen hatırlatmalar:** " + " · ".join(f"{r['title']} ({reminder_time(r):%d.%m %H:%M})" for r in due[:5]))
    if today:
        st.warning("📅 **Bugün yaklaşan hatırlatmalar:** " + " · ".join(f"{r['title']} ({reminder_time(r):%H:%M})" for r in today[:5]))
    if not due and not today:
        st.success("Bugün için bekleyen hatırlatmanız yok.")
    generation = st.session_state.get("agenda_generation", 0)

    def bump() -> None:
        st.session_state["agenda_generation"] = generation + 1

    labels = {str(r["id"]): f"{_clean(r.get('title'))} — {_clean(r.get('city'))}/{_clean(r.get('district'))}" for r in properties.to_dict("records")} if not properties.empty else {}
    now = now_tr()
    with st.expander("➕ Hatırlatma ekle", expanded=not reminders):
        title = st.text_input("Hatırlatma başlığı *", key=f"rem_title_{generation}", placeholder="Örn. Zeynep Hanım'ı ara, anahtar teslimi…")
        d1, d2, d3 = st.columns(3)
        remind_date = d1.date_input("Tarih", value=now.date(), key=f"rem_date_{generation}")
        remind_time = d2.time_input("Saat", value=(now + timedelta(hours=1)).time().replace(minute=0, second=0, microsecond=0), key=f"rem_time_{generation}")
        linked = d3.selectbox("İlgili ilan (isteğe bağlı)", ["Yok"] + list(labels), format_func=lambda x: x if x == "Yok" else labels[x], key=f"rem_property_{generation}")
        notes = st.text_area("Not", height=70, key=f"rem_notes_{generation}")
        if st.button("Hatırlatmayı kaydet", type="primary", key=f"rem_save_{generation}"):
            if not title.strip():
                st.error("Hatırlatma başlığını yazın.")
            else:
                client.table("reminders").insert({"user_id": user_id, "title": title.strip(), "notes": notes.strip(), "remind_at": datetime.combine(remind_date, remind_time, tzinfo=TZ).isoformat(), "property_id": None if linked == "Yok" else linked}).execute()
                bump()
                st.rerun()
    show = st.radio("Göster", ["Bekleyenler", "Tamamlananlar", "Tümü"], horizontal=True, key="agenda_filter", label_visibility="collapsed")
    visible = [r for r in reminders if (show == "Tümü") or (show == "Tamamlananlar" and r.get("done")) or (show == "Bekleyenler" and not r.get("done"))]
    st.caption(f"{len(visible)} hatırlatma")
    with st.container(height=380, border=False):
        for reminder in visible:
            reminder_id = str(reminder["id"])
            when = reminder_time(reminder)
            icon = "✅" if reminder.get("done") else ("🔴" if when <= now else ("🟡" if when.date() == now.date() else "⏰"))
            with st.expander(f"{icon} {when:%d.%m.%Y %H:%M} — {reminder['title']}"):
                property_id = str(reminder.get("property_id") or "")
                if property_id in labels:
                    st.button(f"🏠 İlanı aç: {labels[property_id]}", key=f"rem_open_{reminder_id}_{generation}", on_click=open_property, args=(property_id,))
                new_title = st.text_input("Başlık", value=reminder["title"], key=f"rem_edit_title_{reminder_id}_{generation}")
                e1, e2 = st.columns(2)
                new_date = e1.date_input("Tarih", value=when.date(), key=f"rem_edit_date_{reminder_id}_{generation}")
                new_time = e2.time_input("Saat", value=when.time().replace(second=0, microsecond=0), key=f"rem_edit_time_{reminder_id}_{generation}")
                new_notes = st.text_area("Not", value=reminder.get("notes") or "", height=70, key=f"rem_edit_notes_{reminder_id}_{generation}")
                b1, b2, b3 = st.columns(3)
                if b1.button("Kaydet", type="primary", key=f"rem_update_{reminder_id}_{generation}", use_container_width=True):
                    client.table("reminders").update({"title": new_title.strip() or reminder["title"], "notes": new_notes.strip(), "remind_at": datetime.combine(new_date, new_time, tzinfo=TZ).isoformat()}).eq("id", reminder_id).execute()
                    bump()
                    st.rerun()
                if b2.button("Tekrar aç" if reminder.get("done") else "Tamamlandı", key=f"rem_done_{reminder_id}_{generation}", use_container_width=True):
                    client.table("reminders").update({"done": not reminder.get("done")}).eq("id", reminder_id).execute()
                    bump()
                    st.rerun()
                if b3.button("Sil", key=f"rem_delete_{reminder_id}_{generation}", use_container_width=True):
                    client.table("reminders").delete().eq("id", reminder_id).execute()
                    bump()
                    st.rerun()


def haversine_km(lat1: float, lng1: float, lat2: float, lng2: float) -> float:
    p1, p2 = math.radians(lat1), math.radians(lat2)
    a = math.sin((p2 - p1) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(math.radians(lng2 - lng1) / 2) ** 2
    return 6371.0 * 2 * math.asin(math.sqrt(a))


@st.cache_data(ttl=3600, show_spinner=False)
def road_route(lat1: float, lng1: float, lat2: float, lng2: float) -> dict | None:
    try:
        response = requests.get(f"https://router.project-osrm.org/route/v1/driving/{lng1},{lat1};{lng2},{lat2}", params={"overview": "false"}, headers=USER_AGENT, timeout=10)
        route = response.json()["routes"][0]
        return {"km": route["distance"] / 1000, "minutes": route["duration"] / 60}
    except Exception:
        return None


def distance_tool(property_row: pd.Series) -> None:
    property_id = str(property_row.id)
    st.markdown("##### 📏 Mesafe ölç")
    city, district = _clean(property_row.get("city")), _clean(property_row.get("district"))
    latitude, longitude = pd.to_numeric(property_row.get("latitude"), errors="coerce"), pd.to_numeric(property_row.get("longitude"), errors="coerce")
    origin = (float(latitude), float(longitude)) if pd.notna(latitude) and pd.notna(longitude) else None
    q1, q2 = st.columns([4, 1], vertical_alignment="bottom")
    query = q1.text_input("Mesafesini öğrenmek istediğiniz adres veya yer", key=f"dist_query_{property_id}", placeholder="Örn. Selçuklu Devlet Hastanesi, otogar, çalıştığınız iş yeri adresi…")
    if q2.button("Mesafeyi ölç", key=f"dist_button_{property_id}", use_container_width=True):
        if not query.strip():
            st.warning("Bir adres veya yer adı yazın.")
        else:
            try:
                start = origin
                if start is None:
                    place = search_location(", ".join(part for part in (_clean(property_row.get("address")), district, city, "Türkiye") if part))
                    start = (float(place["lat"]), float(place["lon"])) if place else None
                target = search_location(f"{query.strip()}, {city}, Türkiye") or search_location(f"{query.strip()}, Türkiye")
            except Exception:
                start, target = None, None
                st.error("Adres arama servisine şu anda ulaşılamadı.")
            if start is None:
                st.warning("İlanın konumu bulunamadı. İlan düzenleme ekranından haritada konumu işaretleyin.")
            elif not target:
                st.warning("Girilen adres bulunamadı; daha ayrıntılı yazmayı deneyin.")
            else:
                end = (float(target["lat"]), float(target["lon"]))
                route = road_route(round(start[0], 6), round(start[1], 6), round(end[0], 6), round(end[1], 6))
                st.session_state[f"dist_{property_id}"] = {"start": start, "end": end, "straight": haversine_km(*start, *end), "route": route, "label": target.get("display_name", query)}
    result = st.session_state.get(f"dist_{property_id}")
    if not result:
        return
    m1, m2, m3 = st.columns(3)
    m1.metric("Kuş uçuşu", f"{result['straight']:.2f} km")
    if result["route"]:
        m2.metric("Araçla yol", f"{result['route']['km']:.1f} km")
        m3.metric("Tahmini süre", f"{result['route']['minutes']:.0f} dk")
    else:
        m2.metric("Araçla yol", "—")
        m3.caption("Yol bilgisi alınamadı; kuş uçuşu mesafe gösteriliyor.")
    st.caption("Hedef: " + str(result["label"]).split(",")[0] + " — " + ", ".join(str(result["label"]).split(",")[1:3]).strip())
    fmap = folium.Map(location=[(result["start"][0] + result["end"][0]) / 2, (result["start"][1] + result["end"][1]) / 2], zoom_start=13)
    folium.Marker(result["start"], tooltip="İlan", icon=folium.Icon(color="red")).add_to(fmap)
    folium.Marker(result["end"], tooltip="Hedef", icon=folium.Icon(color="blue")).add_to(fmap)
    folium.PolyLine([result["start"], result["end"]], color="#1d4ed8", weight=3, dash_array="6").add_to(fmap)
    fmap.fit_bounds([result["start"], result["end"]], padding=(30, 30))
    st_folium(fmap, height=260, use_container_width=True, key=f"dist_map_{property_id}", returned_objects=[])


def _tr_upper(text: str) -> str:
    return text.replace("i", "İ").replace("ı", "I").upper()


def whatsapp_caption(row: pd.Series, phone: str = "", eids: bool = True) -> str:
    def get(name: str) -> str:
        value = _clean(row.get(name))
        return "" if value in ("Belirtilmemiş", "Belirtilmedi") else value

    lines = [f"🔥 {_tr_upper(get('title'))} 🔥"]
    description = " ".join(get("description").split())
    if description:
        lines.append(description if len(description) <= 300 else description[:297].rstrip() + "…")
    lines.append("")
    address = get("address").split(",")[0].strip()
    district, city = get("district"), get("city")
    location = " / ".join(part for part in (district, address) if part)
    lines.append(f"📍 Konum: {location}{' – ' + city if city else ''}")
    gross, net = get("area_m2"), get("net_area_m2")
    if gross and net and net != "0":
        lines.append(f"📐 Metrekare: {net} m² Net / {gross} m² Brüt")
    elif gross:
        lines.append(f"📐 Metrekare: {gross} m²")
    if get("rooms"):
        kitchen = {"Kapalı": " (Kapalı Mutfak)", "Açık (Amerikan)": " (Açık Mutfak)"}.get(get("kitchen"), "")
        lines.append(f"🏠 Oda Sayısı: {get('rooms')}{kitchen}")
    if get("heating"):
        heating = get("heating")
        lines.append(f"♨️ Isıtma: {heating + ' Isıtma' if heating in ('Merkezi', 'Yerden ısıtma') and 'ısıtma' not in heating.lower() else heating}")
    if get("bathroom_count"):
        lines.append(f"🚿 Banyo: {get('bathroom_count')}")
    if get("parking") and get("parking") != "Yok":
        lines.append(f"🚗 Otopark: {'Sitede ' if row.get('within_complex') else ''}{get('parking')} Otopark")
    if get("title_deed_status"):
        loan = " (Krediye Uygun)" if row.get("loan_eligible") else ""
        lines.append(f"🚪 Tapu Durumu: {get('title_deed_status')}{loan}")
    lines.append("Detaylı bilgi ve randevu için hemen iletişime geçin:")
    if eids:
        lines += ["", "✅ EİDS Onaylı"]
    price = money_text(row.get("price"))
    lines += ["", f"🚨🚨🚨 *{price} TL{' / Ay' if get('deal_type') == 'Kiralık' else ''}*"]
    if phone.strip():
        lines += ["", f"📞 *{phone.strip()}*"]
    return "\n".join(lines)


def _hex_rgb(value: str) -> tuple[int, int, int]:
    value = (value or "#F26B0F").lstrip("#")
    try:
        return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return (242, 107, 15)


def _shade(color: tuple[int, int, int], factor: float) -> tuple[int, int, int]:
    return tuple(max(0, min(255, int(c * factor))) for c in color)


def _draw_whatsapp_icon(draw, cx: float, cy: float, r: float) -> None:
    white, green = (255, 255, 255, 255), (37, 211, 102, 255)
    draw.ellipse([cx - r, cy - r, cx + r, cy + r], fill=white)
    draw.polygon([(cx - r * 0.98, cy + r * 0.98), (cx - r * 0.25, cy + r * 0.8), (cx - r * 0.92, cy + r * 0.2)], fill=white)
    inner = r * 0.84
    draw.ellipse([cx - inner, cy - inner, cx + inner, cy + inner], fill=green)
    draw.polygon([(cx - r * 0.86, cy + r * 0.86), (cx - r * 0.3, cy + r * 0.72), (cx - r * 0.8, cy + r * 0.3)], fill=green)
    hx, hy = cx + r * 0.04, cy - r * 0.04
    outer, inner_r, start, end, steps = r * 0.52, r * 0.30, 85, 215, 40
    points = [(hx + outer * math.cos(math.radians(start + (end - start) * i / steps)), hy + outer * math.sin(math.radians(start + (end - start) * i / steps))) for i in range(steps + 1)]
    points += [(hx + inner_r * math.cos(math.radians(start + (end - start) * i / steps)), hy + inner_r * math.sin(math.radians(start + (end - start) * i / steps))) for i in range(steps, -1, -1)]
    draw.polygon(points, fill=white)
    middle = (outer + inner_r) / 2
    for angle in (start, end):
        ex, ey = hx + middle * math.cos(math.radians(angle)), hy + middle * math.sin(math.radians(angle))
        draw.ellipse([ex - r * 0.19, ey - r * 0.19, ex + r * 0.19, ey + r * 0.19], fill=white)


def _draw_note(canvas, text: str, color: str) -> None:
    from PIL import Image, ImageDraw

    width, height = canvas.size
    note_w, note_h, margin = int(width * 0.30), int(height * 0.22), 50
    layer = Image.new("RGBA", (note_w + 2 * margin, note_h + 2 * margin), (0, 0, 0, 0))
    draw = ImageDraw.Draw(layer)
    paper = _hex_rgb(color)
    draw.rectangle([margin + 8, margin + 9, margin + note_w + 8, margin + note_h + 9], fill=(0, 0, 0, 60))
    for y in range(note_h):
        draw.line([(margin, margin + y), (margin + note_w, margin + y)], fill=_shade(paper, 1.07 - 0.13 * y / note_h) + (255,))

    def tape(cx: int, cy: int, angle: int) -> None:
        strip = Image.new("RGBA", (int(note_w * 0.26), int(note_h * 0.13)), (238, 226, 200, 215)).rotate(angle, expand=True, resample=Image.BICUBIC)
        layer.paste(strip, (cx - strip.width // 2, cy - strip.height // 2), strip)

    tape(margin + int(note_w * 0.05), margin + int(note_h * 0.04), 38)
    tape(margin + int(note_w * 0.95), margin + int(note_h * 0.04), -38)
    pad, icon_r = int(note_w * 0.07), int(note_h * 0.17)
    icon_cx, icon_cy = margin + pad + icon_r, margin + pad + icon_r + int(note_h * 0.03)
    _draw_whatsapp_icon(draw, icon_cx, icon_cy, icon_r)
    words = text.split()
    size_px = int(note_h * 0.2)
    while True:
        font = _font(size_px, note=True)
        first_width, other_width = note_w - (2 * icon_r + 3 * pad), note_w - 2 * pad
        lines, current, limit = [], "", first_width
        for word in words:
            trial = f"{current} {word}".strip()
            if draw.textlength(trial, font=font) <= limit:
                current = trial
            else:
                if current:
                    lines.append(current)
                    limit = other_width
                current = word
        lines.append(current)
        line_h = int(size_px * 1.2)
        bottom = margin + pad + 2 * icon_r + int(pad * 0.8) + max(len(lines) - 1, 0) * line_h + line_h
        if bottom <= margin + note_h - pad // 2 or size_px <= 14:
            break
        size_px -= 2
    ink, outline = (255, 230, 110), _shade(paper, 0.5)
    for index, line in enumerate(lines):
        if index == 0:
            x, y = icon_cx + icon_r + pad, icon_cy - int(size_px * 0.6)
        else:
            x, y = margin + pad, margin + pad + 2 * icon_r + int(pad * 0.8) + (index - 1) * line_h
        draw.text((x, y), line, font=font, fill=ink, stroke_width=2, stroke_fill=outline)
    layer = layer.rotate(4, expand=True, resample=Image.BICUBIC)
    canvas.paste(layer, (int(width * 0.005), int(height * 0.005)), layer)


def make_collage(images: list, label: str, style: str = "ribbon", color: str = "#F26B0F", size: tuple[int, int] = (1200, 900)):
    from PIL import Image, ImageDraw, ImageOps

    width, height = size
    gap = 5
    canvas = Image.new("RGB", size, "white")
    half_w, half_h = width // 2, height // 2
    layouts = {
        1: [(0, 0, width, height)],
        2: [(0, 0, half_w - gap, height), (half_w + gap, 0, width, height)],
        3: [(0, 0, half_w - gap, height), (half_w + gap, 0, width, half_h - gap), (half_w + gap, half_h + gap, width, height)],
        4: [(0, 0, half_w - gap, half_h - gap), (half_w + gap, 0, width, half_h - gap), (0, half_h + gap, half_w - gap, height), (half_w + gap, half_h + gap, width, height)],
    }
    count = min(len(images), 4)
    if count == 0:
        ImageDraw.Draw(canvas).rectangle([0, 0, width, height], fill="#f2f4f7")
    for box, image in zip(layouts.get(count, []), images[:4]):
        canvas.paste(ImageOps.fit(image, (box[2] - box[0], box[3] - box[1]), method=Image.LANCZOS), (box[0], box[1]))
    words = label.replace("\n", " ").split()
    if not words or style == "none":
        return canvas
    if style == "note":
        _draw_note(canvas, " ".join(words), color)
        return canvas
    s = width / 1448
    base = _hex_rgb(color)
    overlay = Image.new("RGBA", size, (0, 0, 0, 0))
    layer = ImageDraw.Draw(overlay)
    ribbon_w, top = int(560 * s), int(30 * s)
    points = [(0, top + int(40 * s)), (ribbon_w, top), (ribbon_w - int(35 * s), top + int(205 * s)), (0, top + int(245 * s))]
    layer.polygon([(x + int(8 * s), y + int(10 * s)) for x, y in points], fill=(0, 0, 0, 110))
    layer.polygon(points, fill=base + (255,), outline=_shade(base, 0.55) + (255,))
    layer.polygon([(0, top + int(40 * s)), (ribbon_w, top), (ribbon_w - int(8 * s), top + int(60 * s)), (0, top + int(95 * s))], fill=(255, 255, 255, 50))
    canvas = Image.alpha_composite(canvas.convert("RGBA"), overlay).convert("RGB")
    draw = ImageDraw.Draw(canvas)
    first, last = (" ".join(words[:-1]), words[-1]) if len(words) > 1 else ("", words[0])

    def fit(text: str, start: int, max_width: int):
        size_px = start
        while size_px > 16 and draw.textlength(text, font=_font(size_px, True)) > max_width:
            size_px -= 2
        return _font(size_px, True)

    stroke = _shade(base, 0.4)
    max_width = ribbon_w - int(110 * s)
    if first:
        draw.text((int(45 * s), top + int(52 * s)), first, font=fit(first, int(54 * s), max_width), fill="white", stroke_width=max(2, int(3 * s)), stroke_fill=stroke)
    draw.text((int(55 * s), top + int(100 * s) if first else top + int(70 * s)), last, font=fit(last, int(100 * s), max_width), fill=(255, 235, 120), stroke_width=max(3, int(5 * s)), stroke_fill=stroke)
    return canvas


def _jpeg_b64(image, max_side: int = 1280, quality: int = 82) -> str:
    image = image.convert("RGB")
    image.thumbnail((max_side, max_side))
    buffer = io.BytesIO()
    image.save(buffer, "JPEG", quality=quality)
    return base64.b64encode(buffer.getvalue()).decode("ascii")


WHATSAPP_HTML = """<!doctype html><html><head><meta charset="utf-8"><style>
body{margin:0;font-family:system-ui,sans-serif}
button{width:100%;padding:12px;border:0;border-radius:10px;font-size:15px;font-weight:700;cursor:pointer;margin-bottom:8px}
.main{background:#25D366;color:#fff}
.row{display:flex;gap:8px}
.row button{background:#e8f5e9;color:#14532d;padding:10px;font-size:13px}
#msg{margin-top:2px;font-size:13px;color:#475467}
</style></head><body>
<button id="seq" class="main" type="button">📲 Resim + mesaj gönder (önce resim, sonra mesaj)</button>
<div class="row"><button id="files" type="button">Yalnızca resim</button><button id="textonly" type="button">Yalnızca mesaj</button><button id="copy" type="button">Metni kopyala</button></div>
<div id="msg"></div>
<script>
const files64 = __FILES__; const text = __TEXT__;
const msg = document.getElementById("msg"); const seq = document.getElementById("seq");
const seqLabel = seq.textContent; let step = 0;
if (!files64.length) {
  seq.style.display = "none"; document.getElementById("files").style.display = "none";
  document.getElementById("textonly").textContent = "Mesajı paylaş";
}
function makeFiles() {
  return files64.map((b64, i) => {
    const bin = atob(b64); const bytes = new Uint8Array(bin.length);
    for (let k = 0; k < bin.length; k++) bytes[k] = bin.charCodeAt(k);
    return new File([bytes], "ilan_" + (i + 1) + ".jpg", {type: "image/jpeg"});
  });
}
async function tryShare(data) {
  const navs = [navigator];
  try { navs.push(window.parent.navigator); } catch (e) {}
  for (const nav of navs) {
    try {
      if (nav.share && (!data.files || (nav.canShare && nav.canShare(data)))) { await nav.share(data); return "ok"; }
    } catch (e) { if (e && e.name === "AbortError") return "abort"; }
  }
  return "no";
}
async function copyText() {
  try { await navigator.clipboard.writeText(text); return true; } catch (e) {}
  try { await window.parent.navigator.clipboard.writeText(text); return true; } catch (e) {}
  const area = document.createElement("textarea"); area.value = text; document.body.appendChild(area); area.select();
  let ok = false; try { ok = document.execCommand("copy"); } catch (e) {}
  area.remove(); return ok;
}
function download(files) {
  files.forEach((file) => {
    const a = document.createElement("a"); a.href = URL.createObjectURL(file); a.download = file.name;
    document.body.appendChild(a); a.click(); a.remove();
  });
}
function openWhatsApp() { window.open("https://wa.me/?text=" + encodeURIComponent(text), "_blank"); }
async function shareText() {
  const result = await tryShare({text: text});
  if (result !== "no") return result;
  const copied = await copyText();
  msg.textContent = copied ? "Mesaj kopyalandı; WhatsApp'ta yapıştırabilirsiniz." : "";
  openWhatsApp();
  return "fallback";
}
seq.onclick = async () => {
  if (step === 0) {
    const files = makeFiles();
    const result = await tryShare({files: files});
    if (result === "ok") { step = 1; seq.textContent = "② Şimdi mesajı gönder ➜"; msg.textContent = "Resim hazır. Şimdi aynı düğmeye tekrar dokunup mesajı gönderin."; return; }
    if (result === "abort") return;
    download(files); await copyText();
    msg.textContent = "Resimler indirildi, mesaj kopyalandı. WhatsApp'ta önce resimleri ekleyin, sonra mesajı yapıştırın.";
  } else {
    const result = await shareText();
    if (result === "ok" || result === "fallback") { step = 0; seq.textContent = seqLabel; if (result === "ok") msg.textContent = "Tamamlandı."; }
  }
};
document.getElementById("files").onclick = async () => {
  const files = makeFiles();
  if ((await tryShare({files: files})) === "no") { download(files); msg.textContent = "Resimler indirildi; WhatsApp'ta sohbeti seçip ekleyin."; }
};
document.getElementById("textonly").onclick = shareText;
document.getElementById("copy").onclick = async () => {
  msg.textContent = (await copyText()) ? "Mesaj kopyalandı." : "Kopyalanamadı; metni yukarıdaki kutudan elle kopyalayın.";
};
</script></body></html>"""


def whatsapp_block(row: pd.Series, facts: list[tuple[str, str]], photos: list[dict], office_name: str) -> None:
    property_id = str(row.id)
    st.markdown("##### 📲 WhatsApp'ta paylaş")
    mode = st.radio("Paylaşım biçimi", ["Kolaj + açıklama", "İlan sayfası + açıklama", "Yalnızca açıklama"], horizontal=True, key=f"wa_mode_{property_id}")
    c1, c2 = st.columns(2)
    default_phone = st.session_state.get("staff_phones", {}).get(str(row.get("owner_id")), "")
    phone = c1.text_input("Telefon (açıklamada görünür)", value=default_phone, key=f"wa_phone_{property_id}")
    eids = c2.checkbox("✅ EİDS Onaylı satırı ekle", value=True, key=f"wa_eids_{property_id}")
    label_style, label_text, color, singles = "none", "", "#F26B0F", False
    if mode == "Kolaj + açıklama":
        l1, l2, l3 = st.columns([2, 3, 1.2], vertical_alignment="bottom")
        style_choice = l1.selectbox("Kolaj etiketi", ["Şerit etiket", "WhatsApp not kağıdı", "Etiket yok"], key=f"wa_label_style_{property_id}")
        defaults = {"Şerit etiket": ("ribbon", "KANALA ÖZEL FİYAT"), "WhatsApp not kağıdı": ("note", "Sadece WhatsApp Kanalına Özel")}
        if style_choice in defaults:
            label_style = defaults[style_choice][0]
            label_text = l2.text_input("Etiket yazısı", value=defaults[style_choice][1], key=f"wa_label_text_{property_id}_{label_style}")
            color = l3.color_picker("Etiket rengi", "#F26B0F", key=f"wa_color_{property_id}")
        singles = st.checkbox("Fotoğrafları tek tek de ekle (en fazla 4)", value=False, key=f"wa_singles_{property_id}")
    generated = whatsapp_caption(row, phone, eids)
    caption = st.text_area("Paylaşım metni (göndermeden önce düzenleyebilirsiniz)", value=generated, height=260, key=f"wa_caption_{property_id}_{abs(hash(generated))}")
    if mode == "Yalnızca açıklama":
        embed_html(WHATSAPP_HTML.replace("__FILES__", "[]").replace("__TEXT__", json.dumps(caption).replace("</", "<\\/")), 70)
        return
    signature = f"{mode}|{label_style}|{label_text}|{color}|{singles}|{hash(tuple(facts))}|{'|'.join(p['name'] for p in photos[:4])}"
    state_key = f"wa_files_{property_id}"
    prepared = st.session_state.get(state_key)
    if prepared and prepared.get("sig") != signature:
        prepared = None
    if not prepared:
        if st.button("Paylaşım resimlerini hazırla", type="primary", key=f"wa_prepare_{property_id}", use_container_width=True):
            try:
                with st.spinner("Paylaşım resimleri hazırlanıyor…"):
                    images = _fetch_images(photos, 4)
                    if mode == "Kolaj + açıklama":
                        files = [_jpeg_b64(make_collage(images, label_text, label_style, color), 1200)]
                        if singles:
                            files += [_jpeg_b64(image) for image in images]
                    else:
                        page = render_listing_image(facts, images, office_name)
                        files = [_jpeg_b64(page.resize((1000, int(page.height * 1000 / page.width))), 1800)]
                st.session_state[state_key] = {"sig": signature, "files": files}
                st.rerun()
            except Exception as exc:
                st.error("Paylaşım resmi hazırlanamadı.")
                st.caption(str(exc))
        st.caption("Önce paylaşım resmi hazırlanır. Ardından tek düğmeyle önce resim, sonra mesaj WhatsApp'ta gönderilir.")
    else:
        st.image(base64.b64decode(prepared["files"][0]), width=300, caption="Paylaşılacak ana resim")
        embed_html(WHATSAPP_HTML.replace("__FILES__", json.dumps(prepared["files"])).replace("__TEXT__", json.dumps(caption).replace("</", "<\\/")), 190)


def avatar_bytes(client: Client, user_id: str) -> bytes | None:
    store = thumb_store()
    key = f"avatar:{user_id}"
    cached = store.get(key)
    if cached and cached[0] > time.time():
        return cached[1]
    data = None
    try:
        bucket = client.storage.from_(PHOTO_BUCKET)
        objects = bucket.list(f"avatars/{user_id}", {"limit": 5}) or []
        first = next((o for o in objects if o.get("id")), None)
        data = bucket.download(f"avatars/{user_id}/{first['name']}") if first else None
    except Exception:
        data = None
    store[key] = (time.time() + 900, data)
    return data


def save_avatar(client: Client, user_id: str, upload) -> None:
    from PIL import Image, ImageOps

    image = ImageOps.exif_transpose(Image.open(io.BytesIO(upload.getvalue()))).convert("RGB")
    image = ImageOps.fit(image, (400, 400), method=Image.LANCZOS)
    buffer = io.BytesIO()
    image.save(buffer, "JPEG", quality=85)
    bucket = service_client(client).storage.from_(PHOTO_BUCKET)
    bucket.upload(f"avatars/{user_id}/avatar.jpg", buffer.getvalue(), file_options={"content-type": "image/jpeg", "upsert": "true"})
    thumb_store().pop(f"avatar:{user_id}", None)


def remove_avatar(client: Client, user_id: str) -> None:
    service_client(client).storage.from_(PHOTO_BUCKET).remove([f"avatars/{user_id}/avatar.jpg"])
    thumb_store().pop(f"avatar:{user_id}", None)


def show_avatar(client: Client, user_id: str, name: str, width: int = 96) -> None:
    data = avatar_bytes(client, user_id)
    if data:
        st.image(data, width=width)
    else:
        st.markdown(f"<div class='avatar-ph'>{escape((name or '?')[:1].upper())}</div>", unsafe_allow_html=True)


HERO_SECONDS = 7


def _fetch_hero(client: Client, store: dict, property_id: str) -> None:
    from PIL import Image, ImageOps

    key = f"hero:{property_id}"
    try:
        bucket = client.storage.from_(PHOTO_BUCKET)
        objects = bucket.list(property_id, {"limit": 5, "sortBy": {"column": "name", "order": "asc"}}) or []
        first = next((o for o in objects if o.get("id")), None)
        encoded = None
        if first:
            image = ImageOps.exif_transpose(Image.open(io.BytesIO(bucket.download(f"{property_id}/{first['name']}")))).convert("RGB")
            image = ImageOps.fit(image, (1200, 480), method=Image.LANCZOS)
            buffer = io.BytesIO()
            image.save(buffer, "JPEG", quality=80)
            encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
        store[key] = (time.time() + 900, encoded)
    except Exception:
        store[key] = (time.time() + 60, None)


def ensure_heroes(client: Client, property_ids: list[str], budget: float = 7.0) -> None:
    store = thumb_store()
    now = time.time()
    missing = [pid for pid in property_ids if f"hero:{pid}" not in store or store[f"hero:{pid}"][0] < now]
    if not missing:
        return
    pool = ThreadPoolExecutor(max_workers=6)
    try:
        wait_futures([pool.submit(_fetch_hero, client, store, pid) for pid in missing], timeout=budget)
    finally:
        pool.shutdown(wait=False)


def hero_step(delta: int, count: int) -> None:
    st.session_state["hero_index"] = (st.session_state.get("hero_index", 0) + delta) % max(count, 1)
    st.session_state["hero_changed"] = time.time()


def hero_banner(client: Client, properties: pd.DataFrame) -> None:
    if properties.empty or "status" not in properties:
        return
    candidates = properties[properties.status == "Aktif"]
    if "created_at" in candidates:
        candidates = candidates.sort_values("created_at", ascending=False)
    candidates = candidates.head(10).to_dict("records")
    if not candidates:
        return
    ensure_heroes(client, [str(c["id"]) for c in candidates])
    store = thumb_store()
    slides = [(c, (store.get(f"hero:{c['id']}") or (0, None))[1]) for c in candidates]
    slides = ([s for s in slides if s[1]] or slides)[:6]
    fallback = "data:image/svg+xml;base64," + base64.b64encode(DEFAULT_PROPERTY_IMAGE.read_bytes()).decode("ascii") if DEFAULT_PROPERTY_IMAGE.exists() else ""

    def render() -> None:
        count = len(slides)
        if time.time() - st.session_state.get("hero_changed", 0) >= HERO_SECONDS and count > 1:
            hero_step(1, count)
        index = st.session_state.get("hero_index", 0) % count
        item, encoded = slides[index]
        image_url = f"data:image/jpeg;base64,{encoded}" if encoded else fallback
        badge = escape(_tr_upper(f"{_clean(item.get('deal_type'))} {_clean(item.get('property_category')) or 'Daire'}").strip())
        price = int(float(item.get("price") or 0))
        area = _clean(item.get("area_m2"))
        meta = " · ".join(part for part in (_clean(item.get("rooms")), f"{area} m²" if area else "") if part)
        st.markdown(
            f"<div class='hero' style=\"background-image:url('{image_url}')\"><div class='hero-shade'></div><div class='hero-body'>"
            f"<span class='hero-badge'>{badge}</span><div class='hero-title'>{escape(_clean(item.get('title')) or 'İlan')}</div>"
            f"<div class='hero-meta'><span class='hero-price'>{money_text(price) + ' ₺' if price else 'Fiyat sorunuz'}</span>"
            f"📍 {escape(_clean(item.get('district')))}, {escape(_clean(item.get('city')))}{' · ' + escape(meta) if meta else ''}</div></div></div>",
            unsafe_allow_html=True,
        )
        with st.container(key="hero_controls"):
            c1, c2, c3 = st.columns([1, 6, 1], vertical_alignment="center")
            c1.button("‹", key="hero_prev", use_container_width=True, on_click=hero_step, args=(-1, count))
            if c2.button("İlanı aç", key="hero_open", type="primary", use_container_width=True):
                open_property(str(item["id"]))
                st.rerun()
            c3.button("›", key="hero_next", use_container_width=True, on_click=hero_step, args=(1, count))
        if count > 1:
            st.markdown("<div class='hero-dots'>" + " ".join("●" if i == index else "○" for i in range(count)) + "</div>", unsafe_allow_html=True)

    if hasattr(st, "fragment") and len(slides) > 1:
        st.fragment(run_every=HERO_SECONDS)(render)()
    else:
        render()


def select_employee(employee_id: str) -> None:
    st.session_state["summary_employee"] = str(employee_id)


def employee_summary_tab(client: Client, staff: pd.DataFrame, properties: pd.DataFrame, requests_df: pd.DataFrame, is_admin: bool, user_id: str) -> None:
    selected = st.session_state.get("summary_employee")
    me = staff[staff.id.astype(str) == user_id]
    with st.expander("👤 Profilimi düzenle"):
        mine = me.iloc[0] if not me.empty else pd.Series(dtype=object)
        job = st.text_input("Ünvan / görev", value=_clean(mine.get("job_title")), key="profile_job", placeholder="Örn. Gayrimenkul danışmanı")
        phone = st.text_input("Telefon", value=_clean(mine.get("phone")), key="profile_phone")
        bio = st.text_area("Hakkımda", value=_clean(mine.get("bio")), height=100, key="profile_bio", placeholder="Uzmanlık alanlarınız, çalıştığınız bölgeler, kısa tanıtım…")
        p1, p2 = st.columns([1, 3], vertical_alignment="center")
        with p1:
            show_avatar(client, user_id, _clean(mine.get("name")), 96)
        with p2:
            photo = st.file_uploader("Profil fotoğrafı", type=["jpg", "jpeg", "png", "webp"], key="profile_photo")
            q1, q2 = st.columns(2)
            if q1.button("Fotoğrafı kaydet", disabled=not photo, key="profile_photo_save", use_container_width=True):
                try:
                    save_avatar(client, user_id, photo)
                    st.success("Profil fotoğrafınız güncellendi.")
                    st.rerun()
                except Exception as exc:
                    st.error("Fotoğraf yüklenemedi. Streamlit Secrets'ta service_role_key tanımlı olmalı.")
                    st.caption(str(exc))
            if q2.button("Fotoğrafı kaldır", disabled=avatar_bytes(client, user_id) is None, key="profile_photo_remove", use_container_width=True):
                try:
                    remove_avatar(client, user_id)
                    st.rerun()
                except Exception as exc:
                    st.error("Fotoğraf kaldırılamadı.")
                    st.caption(str(exc))
        if st.button("Profili kaydet", type="primary", key="profile_save"):
            try:
                service_client(client).table("profiles").update({"job_title": job.strip(), "phone": phone.strip(), "bio": bio.strip()}).eq("id", user_id).execute()
                st.success("Profiliniz güncellendi.")
                st.rerun()
            except Exception as exc:
                feature_setup_message("Profil bilgileri", exc)
    st.caption("Bir çalışanın adına dokunun; profili ve ilanları aşağıda görünür, ilana dokununca ilan açılır.")
    columns = st.columns(3)
    for index, person in enumerate(staff.to_dict("records")):
        person_id = str(person["id"])
        owned = properties[properties.owner_id.astype(str) == person_id] if not properties.empty else pd.DataFrame()
        asked = requests_df[requests_df.owner_id.astype(str) == person_id] if not requests_df.empty else pd.DataFrame()
        active = int((owned.status == "Aktif").sum()) if not owned.empty else 0
        new_asked = int((asked.status == "Yeni").sum()) if not asked.empty else 0
        label = f"👤 {person['name']} — {len(owned)} daire · {active} aktif · {len(asked)} talep ({new_asked} yeni)"
        columns[index % 3].button(label, key=f"employee_{person_id}", use_container_width=True, type="primary" if selected == person_id else "secondary", on_click=select_employee, args=(person_id,))
    if selected:
        person = staff[staff.id.astype(str) == selected]
        if not person.empty:
            row = person.iloc[0]
            lines = [f"<b style='font-size:1.1rem'>{escape(str(row['name']))}</b> <span style='color:#667085'>· {escape(_clean(row.get('job_title')) or 'Çalışan')}</span>"]
            if _clean(row.get("phone")):
                lines.append(f"📞 {escape(_clean(row.get('phone')))}")
            lines.append(escape(_clean(row.get("bio")) or "Bu çalışan henüz kendisi hakkında bilgi eklememiş.").replace("\n", "<br>"))
            card_img, card_text = st.columns([1, 5], vertical_alignment="center")
            with card_img:
                show_avatar(client, selected, str(row["name"]), 96)
            with card_text:
                st.markdown("<div class='price-card'>" + "<br>".join(lines) + "</div>", unsafe_allow_html=True)
            if selected != user_id:
                with st.expander(f"✉️ {row['name']} kişisine mesaj gönder"):
                    compose_message_form(client, user_id, staff, f"profile_{selected}", default_recipient=selected, lock_recipient=True)
            owned = properties[properties.owner_id.astype(str) == selected] if not properties.empty else pd.DataFrame()
            if owned.empty:
                st.info("Bu çalışanın ilanı yok.")
            else:
                st.markdown(f"**{row['name']} — {len(owned)} ilan**")
                render_property_rows(client, owned.to_dict("records"), "emp", height=340)
    if not is_admin:
        st.caption("Çalışan listesi ofis yöneticisi tarafından yönetilir.")


# ---------- detaylı arama ----------
def reset_advanced_search() -> None:
    st.session_state["adv_generation"] = st.session_state.get("adv_generation", 0) + 1


SORT_OPTIONS = {
    "En yeni eklenen": ("created_at", False),
    "En eski eklenen": ("created_at", True),
    "Fiyat (düşükten yükseğe)": ("_price", True),
    "Fiyat (yüksekten düşüğe)": ("_price", False),
    "Brüt m² (büyükten küçüğe)": ("_area", False),
    "Brüt m² (küçükten büyüğe)": ("_area", True),
    "Başlık (A-Z)": ("_title", True),
}


def load_saved_searches(client: Client, user_id: str) -> list[dict]:
    return client.table("saved_searches").select("*").eq("user_id", user_id).order("created_at", desc=True).execute().data or []


def apply_saved_search(filters: dict) -> None:
    new_generation = st.session_state.get("adv_generation", 0) + 1
    for name, value in (filters or {}).items():
        st.session_state[f"adv_{name}_{new_generation}"] = value
    st.session_state["adv_generation"] = new_generation


def map_results(client: Client, frame: pd.DataFrame, favorites: set[str]) -> None:
    latitude = pd.to_numeric(frame.get("latitude"), errors="coerce")
    longitude = pd.to_numeric(frame.get("longitude"), errors="coerce")
    geo = frame[latitude.notna() & longitude.notna()].copy()
    geo["_lat"], geo["_lng"] = latitude[geo.index], longitude[geo.index]
    missing = len(frame) - len(geo)
    c1, c2, c3 = st.columns([3, 2, 1], vertical_alignment="center")
    use_radius = c1.checkbox("📍 Merkez noktasına göre ara (haritaya dokunarak merkez seçin)", key="adv_use_radius")
    radius = c2.slider("Yarıçap (km)", 0.5, 30.0, 5.0, 0.5, key="adv_radius", disabled=not use_radius)
    center = st.session_state.get("adv_map_center")
    c3.button("Merkezi temizle", key="adv_clear_center", disabled=not center, use_container_width=True, on_click=lambda: st.session_state.pop("adv_map_center", None))
    if use_radius and center and not geo.empty:
        geo = geo[geo.apply(lambda r: haversine_km(center[0], center[1], r["_lat"], r["_lng"]) <= radius, axis=1)]
    st.caption(f"Haritada {len(geo)} ilan" + (f" · {missing} ilanın konumu işaretlenmemiş" if missing else "") + (f" · merkeze {radius:g} km içinde" if use_radius and center else ""))
    if center:
        view_center, zoom = center, 12
    elif not geo.empty:
        view_center, zoom = (float(geo["_lat"].mean()), float(geo["_lng"].mean())), 11
    else:
        view_center, zoom = TURKEY_CENTER, 6
    fmap = folium.Map(location=view_center, zoom_start=zoom, control_scale=True)
    for item in geo.to_dict("records"):
        label = f"{_clean(item.get('title'))} — {money_text(item.get('price'))} ₺ · {_clean(item.get('rooms'))} [{item['id']}]"
        color = "orange" if str(item["id"]) in favorites else ("blue" if _clean(item.get("deal_type")) == "Satılık" else "green")
        folium.Marker([item["_lat"], item["_lng"]], tooltip=label, icon=folium.Icon(color=color)).add_to(fmap)
    if center:
        folium.Marker(center, tooltip="Arama merkezi", icon=folium.Icon(color="red", icon="star")).add_to(fmap)
        if use_radius:
            folium.Circle(center, radius=radius * 1000, color="#1d4ed8", fill=True, fill_opacity=0.08).add_to(fmap)
    if not center and len(geo) > 1:
        fmap.fit_bounds([[float(geo["_lat"].min()), float(geo["_lng"].min())], [float(geo["_lat"].max()), float(geo["_lng"].max())]], padding=(30, 30))
    out = st_folium(fmap, height=420, use_container_width=True, key="adv_map", returned_objects=["last_clicked", "last_object_clicked_tooltip"])
    clicked = (out or {}).get("last_clicked")
    if use_radius and clicked and clicked != st.session_state.get("adv_last_click"):
        st.session_state["adv_last_click"] = clicked
        st.session_state["adv_map_center"] = (clicked["lat"], clicked["lng"])
        st.rerun()
    tooltip = (out or {}).get("last_object_clicked_tooltip")
    found = re.search(r"\[([^\]]+)\]\s*$", tooltip or "")
    if found:
        match = geo[geo.id.astype(str) == found.group(1)]
        if not match.empty:
            item = match.iloc[0]
            b1, b2 = st.columns([4, 1], vertical_alignment="center")
            b1.markdown(f"**{_clean(item.get('title'))}** · {money_text(item.get('price'))} ₺ · {_clean(item.get('city'))}/{_clean(item.get('district'))} · {_clean(item.get('rooms'))}")
            b2.button("İlanı aç", key="adv_map_open", type="primary", use_container_width=True, on_click=open_property, args=(str(item["id"]),))


def advanced_search_tab(client: Client, properties: pd.DataFrame, user_id: str, favorites: set[str]) -> None:
    if properties.empty:
        st.info("Henüz daire kaydı yok.")
        return
    generation = st.session_state.get("adv_generation", 0)

    def K(name: str) -> str:
        return f"adv_{name}_{generation}"

    def options(column: str) -> list[str]:
        if column not in properties:
            return []
        return sorted({_clean(v) for v in properties[column].dropna().unique() if _clean(v)})

    def multi(col, label: str, name: str, choices: list[str]):
        current = st.session_state.get(K(name))
        if isinstance(current, list):
            st.session_state[K(name)] = [v for v in current if v in choices]
        return col.multiselect(label, choices, key=K(name))

    if st.session_state.get(K("sort")) not in SORT_OPTIONS:
        st.session_state.pop(K("sort"), None)
    with st.expander("🔎 Filtreler ve sıralama", expanded=True):
        r1 = st.columns([2, 1, 1, 1])
        text = r1[0].text_input("Anahtar kelime", key=K("text"), placeholder="Başlık, adres, site adı, açıklama")
        deals = multi(r1[1], "İşlem", "deal", ["Kiralık", "Satılık"])
        categories = multi(r1[2], "Emlak tipi", "category", options("property_category"))
        statuses = multi(r1[3], "Durum", "status", options("status"))
        r2 = st.columns(4)
        cities = multi(r2[0], "İl", "city", options("city"))
        pool = properties[properties.city.isin(cities)] if cities else properties
        districts = multi(r2[1], "İlçe", "district", sorted({_clean(v) for v in pool.district.dropna().unique() if _clean(v)}))
        rooms = multi(r2[2], "Oda sayısı", "rooms", options("rooms"))
        owners = multi(r2[3], "Sorumlu", "owner", options("sorumlu"))
        r3 = st.columns(4)
        price_min = r3[0].number_input("En düşük fiyat (₺)", min_value=0, step=10000, key=K("price_min"))
        price_max = r3[1].number_input("En yüksek fiyat (₺, 0 = sınırsız)", min_value=0, step=10000, key=K("price_max"))
        area_min = r3[2].number_input("En az brüt m²", min_value=0, step=5, key=K("area_min"))
        area_max = r3[3].number_input("En çok brüt m² (0 = sınırsız)", min_value=0, step=5, key=K("area_max"))
        r4 = st.columns(4)
        ages = multi(r4[0], "Bina yaşı", "age", options("building_age"))
        floors = multi(r4[1], "Bulunduğu kat", "floor", options("floor"))
        heating = multi(r4[2], "Isıtma", "heating", options("heating"))
        parking = multi(r4[3], "Otopark", "parking", options("parking"))
        r5 = st.columns(4)
        usage = multi(r5[0], "Kullanım durumu", "usage", options("usage_status"))
        deeds = multi(r5[1], "Tapu durumu", "deed", options("title_deed_status"))
        sort_choice = r5[2].selectbox("Sıralama", list(SORT_OPTIONS), key=K("sort"))
        r5[3].button("Filtreleri temizle", on_click=reset_advanced_search, use_container_width=True, key=K("reset"))
        r6 = st.columns(6)
        flag_names = {"balcony": "balcony", "elevator": "elevator", "furnished": "furnished", "within_complex": "complex", "loan_eligible": "loan", "swap_allowed": "swap"}
        flag_labels = {"balcony": "Balkon", "elevator": "Asansör", "furnished": "Eşyalı", "within_complex": "Site içinde", "loan_eligible": "Krediye uygun", "swap_allowed": "Takas"}
        flags = {column: r6[index].checkbox(flag_labels[column], key=K(flag_names[column])) for index, column in enumerate(flag_names)}
        filters = {
            "text": text, "deal": deals, "category": categories, "status": statuses, "city": cities, "district": districts, "rooms": rooms, "owner": owners,
            "price_min": price_min, "price_max": price_max, "area_min": area_min, "area_max": area_max, "age": ages, "floor": floors, "heating": heating,
            "parking": parking, "usage": usage, "deed": deeds, "sort": sort_choice, **{flag_names[column]: value for column, value in flags.items()},
        }
        st.divider()
        n1, n2 = st.columns([3, 1], vertical_alignment="bottom")
        search_name = n1.text_input("Bu aramayı kaydet", key="saved_search_name", placeholder="Arama adı yazın, örn. Meram 3+1 satılık")
        if n2.button("Aramayı kaydet", type="primary", key="saved_search_save", use_container_width=True):
            if not search_name.strip():
                st.error("Aramaya bir ad verin.")
            else:
                try:
                    client.table("saved_searches").insert({"user_id": user_id, "name": search_name.strip(), "filters": filters}).execute()
                    st.success("Arama kaydedildi; aşağıdaki “Kayıtlı aramalarım” bölümünden tekrar kullanabilirsiniz.")
                except Exception as exc:
                    feature_setup_message("Kayıtlı aramalar", exc)

    with st.expander("💾 Kayıtlı aramalarım"):
        try:
            saved = load_saved_searches(client, user_id)
        except Exception as exc:
            feature_setup_message("Kayıtlı aramalar", exc)
            saved = None
        if saved is not None:
            if not saved:
                st.caption("Henüz kayıtlı aramanız yok. Filtreleri ayarlayıp yukarıdan kaydedebilirsiniz.")
            else:
                labels = {str(s["id"]): s["name"] for s in saved}
                chosen = st.selectbox("Kayıtlı arama", list(labels), format_func=labels.get, key="saved_search_choice")
                item = next(s for s in saved if str(s["id"]) == chosen)
                b1, b2, b3 = st.columns(3)
                b1.button("Uygula", type="primary", key="saved_search_apply", use_container_width=True, on_click=apply_saved_search, args=(item.get("filters") or {},))
                edit_open = b2.toggle("Düzenle", key=f"saved_search_edit_{chosen}")
                delete_open = b3.toggle("Sil", key=f"saved_search_del_{chosen}")
                if edit_open:
                    new_name = st.text_input("Arama adı", value=item["name"], key=f"saved_search_newname_{chosen}")
                    replace_filters = st.checkbox("Kayıtlı filtreleri şu anki filtrelerle değiştir", key=f"saved_search_replace_{chosen}")
                    if st.button("Değişiklikleri kaydet", key=f"saved_search_update_{chosen}", use_container_width=True):
                        values = {"name": new_name.strip() or item["name"], "updated_at": now_tr().isoformat()}
                        if replace_filters:
                            values["filters"] = filters
                        client.table("saved_searches").update(values).eq("id", chosen).execute()
                        st.rerun()
                if delete_open:
                    st.warning(f"“{item['name']}” kayıtlı araması silinecek.")
                    if st.button("Aramayı sil", key=f"saved_search_delete_{chosen}", use_container_width=True):
                        client.table("saved_searches").delete().eq("id", chosen).execute()
                        st.rerun()


    frame = properties.copy()
    frame["_price"] = pd.to_numeric(frame.get("price"), errors="coerce").fillna(0)
    frame["_area"] = pd.to_numeric(frame.get("area_m2"), errors="coerce").fillna(0)
    frame["_title"] = frame.get("title", pd.Series("", index=frame.index)).fillna("").astype(str).str.lower()
    if text.strip():
        columns = [c for c in ("title", "address", "complex_name", "description", "city", "district") if c in frame]
        blob = frame[columns].fillna("").astype(str).agg(" ".join, axis=1).str.lower()
        frame = frame[blob.str.contains(text.strip().lower(), regex=False)]
    for column, chosen in (("deal_type", deals), ("property_category", categories), ("status", statuses), ("city", cities), ("district", districts),
                           ("rooms", rooms), ("sorumlu", owners), ("building_age", ages), ("floor", floors), ("heating", heating),
                           ("parking", parking), ("usage_status", usage), ("title_deed_status", deeds)):
        if chosen and column in frame:
            frame = frame[frame[column].astype(str).isin(chosen)]
    if price_min:
        frame = frame[frame["_price"] >= price_min]
    if price_max:
        frame = frame[frame["_price"] <= price_max]
    if area_min:
        frame = frame[frame["_area"] >= area_min]
    if area_max:
        frame = frame[frame["_area"] <= area_max]
    for column, wanted in flags.items():
        if wanted and column in frame:
            frame = frame[frame[column].fillna(False).astype(bool)]
    sort_column, ascending = SORT_OPTIONS[sort_choice]
    if sort_column in frame:
        frame = frame.sort_values(sort_column, ascending=ascending, kind="stable")
    average = f" · ortalama fiyat {money_text(frame['_price'].mean())} ₺" if len(frame) else ""
    head1, head2 = st.columns([2, 1], vertical_alignment="center")
    head1.caption(f"{len(frame)} ilan bulundu{average}")
    view_mode = head2.radio("Sonuç görünümü", ["📋 Liste", "🗺️ Harita"], horizontal=True, key="adv_view_mode", label_visibility="collapsed")
    results = frame.drop(columns=["_price", "_area", "_title"])
    if view_mode.startswith("📋"):
        render_property_rows(client, results.to_dict("records"), "adv", height=400, favorites=favorites)
    else:
        map_results(client, results, favorites)


def main() -> None:
    try:
        client = get_client()
    except Exception:
        st.title("🏠 Ofis Portföyü")
        st.error("Supabase bağlantısı ayarlı değil. Kurulum adımları için README.md dosyasına bakın.")
        st.stop()
    if "user" not in st.session_state:
        login(client)
        st.stop()

    # Streamlit reruns the script and creates a fresh Supabase client each time.
    # Restore this browser session's auth tokens so protected profile reads use
    # the authenticated role instead of silently querying as anon.
    try:
        client.auth.set_session(
            st.session_state.supabase_access_token,
            st.session_state.supabase_refresh_token,
        )
        current_session = client.auth.get_session()
        if current_session:
            st.session_state.supabase_access_token = current_session.access_token
            st.session_state.supabase_refresh_token = current_session.refresh_token
    except Exception:
        st.session_state.pop("user", None)
        st.session_state.pop("supabase_access_token", None)
        st.session_state.pop("supabase_refresh_token", None)
        st.warning("Oturum süresi dolmuş. Lütfen yeniden giriş yapın.")
        login(client)
        st.stop()

    user = st.session_state.user
    try:
        staff_rows = rows(client, "profiles")
    except Exception as exc:
        st.error("`profiles` tablosu okunamadı. Supabase API erişim izni veya tablo kurulumu kontrol edilmeli.")
        st.code(str(exc))
        st.stop()
    staff = pd.DataFrame(staff_rows)
    current_id = str(user.id)
    matching_profile = not staff.empty and current_id in staff.get("id", pd.Series(dtype=str)).astype(str).tolist()
    if not matching_profile:
        st.error("Bu Supabase projesindeki `profiles` tablosunda oturum kullanıcınızın UUID’si bulunamadı.")
        st.markdown("### Hızlı eşleştirme kontrolü")
        st.markdown("Aşağıdaki sorguları **Streamlit Secrets’ta tanımlı olan aynı Supabase projesinin** SQL Editor’ında çalıştırın.")
        st.write("Uygulamada giriş yapan kullanıcı UUID’si:")
        st.code(current_id)
        st.write("1. Bu UUID satır döndürüyor mu?")
        st.code(f"select id, email from auth.users where id = '{current_id}';", language="sql")
        st.write("2. Bu UUID için profil var mı?")
        st.code(f"select id, name from public.profiles where id = '{current_id}';", language="sql")
        st.write("3. Profil yoksa yalnızca şu sorguyu çalıştırın:")
        st.code(f"""insert into public.profiles (id, name, is_admin)
values ('{current_id}', 'Adınız Soyadınız', true)
on conflict (id) do update set name = excluded.name;""", language="sql")
        st.caption("2. sorgu satır döndürüyorsa ama uygulama hâlâ bulamıyorsa, uygulamanın Supabase Project URL’si SQL Editor’da açık projenin URL’siyle aynı değildir. `42P07 relation already exists` burada çalıştırılacak sorgu değildir; `create table` içeren tam kurulum dosyasını tekrar çalıştırmayın.")
        st.stop()
    me = staff[staff.id.astype(str) == current_id].iloc[0]
    is_admin = bool(me.get("is_admin", False))
    if not bool(me.get("is_active", True)):
        client.auth.sign_out()
        st.session_state.pop("user", None)
        st.session_state.pop("supabase_access_token", None)
        st.session_state.pop("supabase_refresh_token", None)
        st.error("Hesabınız devre dışı bırakılmış. Ofis yöneticinizle iletişime geçin.")
        st.stop()
    properties = pd.DataFrame(rows(client, "properties"))
    requests = pd.DataFrame(rows(client, "requests"))
    try:
        location_data = get_location_data()
    except Exception:
        location_data = ([], {})
        st.warning("İl/ilçe listesine şu anda ulaşılamadı; adres alanları geçici olarak elle doldurulabilir.")
    staff_names = dict(zip(staff.id, staff.name))
    st.session_state["staff_names"] = {str(k): v for k, v in staff_names.items()}
    st.session_state["staff_phones"] = {str(r["id"]): _clean(r.get("phone")) for r in staff.to_dict("records")}
    settings = get_settings(client)
    office_name = (settings.get("office_name") or DEFAULT_OFFICE_NAME).strip()
    home_title = (settings.get("home_title") or DEFAULT_HOME_TITLE).strip()
    st.session_state["office_name"] = office_name
    favorites, favorites_available = load_favorites(client, current_id)
    try:
        messages_data, messages_available = load_messages(client, current_id), True
    except Exception:
        messages_data, messages_available = [], False
    try:
        reminders_data, reminders_available = load_reminders(client, current_id), True
    except Exception:
        reminders_data, reminders_available = [], False
    unread_messages = sum(1 for m in messages_data if str(m["recipient_id"]) == current_id and not m.get("is_read") and not m.get("deleted_by_recipient"))
    due_reminders, today_reminders = reminder_alerts(reminders_data) if reminders_available else ([], [])
    if not properties.empty:
        properties["sorumlu"] = properties.owner_id.map(staff_names).fillna("Atanmamış")
    if not requests.empty:
        requests["sorumlu"] = requests.owner_id.map(staff_names).fillna("Atanmamış")

    section_names = ["Daireler", "Detaylı arama", "Müşteri talepleri", "Teklifler", "Ajanda", "Mesajlar", "Yeni kayıt", "Çalışan özeti"]
    if is_admin:
        section_names.append("Yönetici")

    with st.sidebar:
        st.button(f"🏠 {office_name} · Ana sayfa", use_container_width=True, key="sidebar_home",
                  on_click=go_to_section, args=("Daireler",))
        st.markdown("### Bölümler")
        badges = {"Ajanda": len(due_reminders) + len(today_reminders), "Mesajlar": unread_messages}
        for name in section_names:
            badge = f" ({badges[name]})" if badges.get(name) else ""
            st.button(name + badge, key=f"nav_{name}", use_container_width=True,
                      on_click=go_to_section, args=(name,))
        st.caption("İlan düzenleme Daireler bölümünde; yedekleme ve geri yükleme Yönetici bölümündedir." if is_admin else "İlan düzenleme Daireler bölümündedir.")
        st.caption(user.email)
        mine = st.toggle("Yalnızca benim kayıtlarım", value=False)
        if st.button("Çıkış yap"):
            client.auth.sign_out()
            st.session_state.pop("user", None)
            st.session_state.pop("supabase_access_token", None)
            st.session_state.pop("supabase_refresh_token", None)
            st.rerun()

    if mine:
        properties = properties[properties.owner_id == user.id] if not properties.empty else properties
        requests = requests[requests.owner_id == user.id] if not requests.empty else requests
    active = int((properties.status == "Aktif").sum()) if not properties.empty else 0
    new_requests = int((requests.status == "Yeni").sum()) if not requests.empty else 0
    agenda_badge = f" ({len(due_reminders) + len(today_reminders)})" if (due_reminders or today_reminders) else ""
    message_badge = f" ({unread_messages})" if unread_messages else ""
    with st.container(key="topbar"):
        nav = st.columns([4, 1, 1, 1.5, 1.2, 1.4], vertical_alignment="center")
        with nav[0]:
            with st.container(key="brand"):
                words = home_title.split()
                first, second = (" ".join(words[:-1]), words[-1]) if len(words) > 1 else (home_title[:-2], home_title[-2:])
                st.markdown(f"<div class='brand'><span class='brand-logo'>🏠</span><span class='brand-text'>{escape(first)} <b>{escape(second)}</b></span></div>", unsafe_allow_html=True)
                st.button(home_title, key="home_title", on_click=go_to_section, args=("Daireler",), help="Daireler bölümüne git")
        nav[1].button("Satılık", key="top_sale", on_click=show_deal, args=("Satılık",), use_container_width=True)
        nav[2].button("Kiralık", key="top_rent", on_click=show_deal, args=("Kiralık",), use_container_width=True)
        nav[3].button("Danışmanlar", key="top_agents", on_click=go_to_section, args=("Çalışan özeti",), use_container_width=True)
        nav[4].button("Ajanda" + agenda_badge, key="top_agenda", on_click=go_to_section, args=("Ajanda",), use_container_width=True)
        nav[5].button("Mesajlar" + message_badge, key="top_messages", on_click=go_to_section, args=("Mesajlar",), use_container_width=True)
    st.markdown(
        f"<div class='summary-strip'><span class='chip'>Daire<b>{len(properties)}</b></span><span class='chip'>Aktif<b>{active}</b></span>"
        f"<span class='chip'>Talep<b>{len(requests)}</b></span><span class='chip'>Yeni talep<b>{new_requests}</b></span></div>",
        unsafe_allow_html=True,
    )
    show_announcements(client)
    if not st.session_state.get("login_alerts_shown"):
        st.session_state["login_alerts_shown"] = True
        if due_reminders or today_reminders:
            st.toast(f"⏰ Ajandanızda {len(due_reminders)} geciken, {len(today_reminders)} bugünkü hatırlatma var.")
        if unread_messages:
            st.toast(f"✉️ {unread_messages} okunmamış mesajınız var.")
    notice = st.session_state.pop("message_notice", None)
    if notice:
        st.toast(notice)
    feature_error = st.session_state.pop("feature_error", None)
    if feature_error:
        feature_setup_message(feature_error[0], Exception(feature_error[1]))
    save_notice = st.session_state.pop("property_save_notice", None)
    if save_notice:
        st.success(f"İlan kaydedildi; {save_notice['uploaded']} fotoğraf yüklendi.")
        for error in save_notice["errors"]:
            st.error(error)
    delete_notice = st.session_state.pop("property_delete_notice", None)
    if delete_notice:
        st.success(delete_notice)

    applied_tab = st.session_state.pop("pending_tab", None)
    if applied_tab not in section_names:
        applied_tab = None
    if applied_tab and TABS_STATE_SUPPORTED:
        st.session_state["main_tabs"] = applied_tab
    if TABS_STATE_SUPPORTED:
        tabs = st.tabs(section_names, key="main_tabs", on_change="rerun")
    else:
        tabs = st.tabs(section_names)
    tab_by_name = dict(zip(section_names, tabs))
    with tab_by_name["Daireler"]:
        selected_property_id = st.session_state.get("selected_property_id")
        detail_row = properties[properties.id.astype(str) == str(selected_property_id)] if selected_property_id and not properties.empty else pd.DataFrame()
        if not detail_row.empty:
            show_property_detail(client, detail_row.iloc[0], current_id, is_admin, staff, location_data, favorites)
        elif properties.empty:
            st.info("Henüz daire kaydı yok.")
        else:
            view_mode = st.radio("Görünüm", ["Tüm ilanlar", "⭐ Favorilerim", "💹 Fiyat değişimleri"], horizontal=True, key="daire_view", label_visibility="collapsed")
            if view_mode == "⭐ Favorilerim":
                if favorites_available:
                    favorites_view(client, current_id, properties, favorites)
                else:
                    feature_setup_message("Favoriler")
            elif view_mode == "💹 Fiyat değişimleri":
                price_changes_view(client, properties)
            else:
                hero_banner(client, properties)
                f1, f2, f3 = st.columns([3, 1.3, 1.3])
                search = f1.text_input("Dairelerde ara", placeholder="İl, ilçe veya başlık")
                deal_filter = f2.selectbox("İşlem türü", ["Tümü", "Kiralık", "Satılık"], key="list_deal_filter")
                status_filter = f3.selectbox("Durum", ["Tümü", "Aktif", "Opsiyonlu", "Kiralandı/Satıldı", "Pasif"])
                view = properties.copy()
                if search:
                    mask = view[["title", "city", "district"]].fillna("").astype(str).apply(lambda col: col.str.contains(search, case=False, regex=False)).any(axis=1)
                    view = view[mask]
                if deal_filter != "Tümü":
                    view = view[view.deal_type == deal_filter]
                if status_filter != "Tümü":
                    view = view[view.status == status_filter]
                st.caption(f"{len(view)} ilan · detay için satıra dokunun")
                render_property_rows(client, view.to_dict("records"), "list", height=360, favorites=favorites)
    with tab_by_name["Detaylı arama"]:
        advanced_search_tab(client, properties, current_id, favorites)
    with tab_by_name["Müşteri talepleri"]:
        requests_tab(client, requests, properties, staff, location_data, current_id, is_admin, favorites)
    with tab_by_name["Teklifler"]:
        offers_tab(client, properties, staff, current_id, is_admin)
    with tab_by_name["Ajanda"]:
        agenda_tab(client, current_id, properties, reminders_data, reminders_available)
    with tab_by_name["Mesajlar"]:
        messages_tab(client, staff, current_id, messages_data, messages_available, properties)
    with tab_by_name["Yeni kayıt"]:
        form1, form2 = st.tabs(["Daire ekle", "Talep ekle"])
        with form1:
            add_property(client, user.id, staff, location_data)
        with form2:
            add_request(client, user.id, staff, location_data)
    with tab_by_name["Çalışan özeti"]:
        employee_summary_tab(client, staff, properties, requests, is_admin, current_id)
    if is_admin:
        with tab_by_name["Yönetici"]:
            try:
                admin_panel(client, get_admin_client(), staff, current_id)
            except Exception as exc:
                st.error("Yönetici işlemleri için Supabase service_role/secret key gerekir. Bu anahtarı Streamlit Cloud Secrets içine ekleyin; GitHub'a koymayın.")
                st.code(str(exc))

    if applied_tab:
        switch_to_tab(applied_tab, click_tab=not TABS_STATE_SUPPORTED)


main()
